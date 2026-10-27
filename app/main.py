"""API y servidor web del Panel de control TI."""

from __future__ import annotations

import datetime as dt
import functools
import io
from urllib.parse import quote, urlparse
import mimetypes
import re
import secrets
import threading
from contextlib import asynccontextmanager
from pathlib import Path

from fastapi import (
    Depends,
    FastAPI,
    File,
    HTTPException,
    Query,
    Request,
    Response,
    UploadFile,
    status,
)
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from sqlalchemy import func, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload, selectinload

from . import schemas
from .auth import (
    COOKIE,
    cerrar_otras_sesiones,
    cerrar_sesion,
    crear_sesion,
    hashear_password,
    poner_cookie,
    quitar_cookie,
    registrar,
    requiere_admin,
    requiere_captura,
    requiere_rehash,
    usuario_actual,
    verificar_password,
    verificar_senuelo,
)
from .auditoria import DIAS_SIN_ACTIVIDAD, PLANTILLAS, construir_items, semana_de
from .database import RUTA_DATOS, SessionLocal, crear_esquema, obtener_sesion
from .models import (
    Activo,
    AlcanceRed,
    Asignacion,
    Auditoria,
    Bitacora,
    Bodega,
    CuentaLicencia,
    Departamento,
    DetalleComputo,
    DetalleCorreo,
    DetalleLicencia,
    DetalleMovil,
    Empleado,
    EstadoAuditoria,
    EstadoDocumento,
    EstadoItem,
    EstadoNota,
    EstatusActivo,
    EstatusEmpleado,
    EvidenciaAuditoria,
    Impresora,
    ItemAuditoria,
    LineaVerificacion,
    Nota,
    Planta,
    PropietarioImpresora,
    RedWifi,
    RolUsuario,
    Sesion,
    AccesoSistema,
    SistemaInterno,
    Sucursal,
    TipoActivo,
    TipoAuditoria,
    Usuario,
    VerificacionNumero,
    a_hora_local,
    ahora_local,
    clave_nombre,
    hoy as hoy_local,
)
from .cifrado import cifrar, cifrar_bytes, descifrar, descifrar_bytes, es_cifrado
from .limite import fallos_por_ip, fallos_por_usuario, peticiones_por_ip, subidas_por_ip
from .seed import PASSWORD_DE_FABRICA, sembrar

ESTATICOS = Path(__file__).parent / "static"
ADJUNTOS = RUTA_DATOS / "adjuntos"
ADJUNTOS.mkdir(parents=True, exist_ok=True)

EXTENSIONES_PERMITIDAS = {".pdf", ".jpg", ".jpeg", ".png", ".webp"}
FIRMAS_ARCHIVO = {
    ".pdf": b"%PDF-",
    ".jpg": b"\xff\xd8\xff",
    ".jpeg": b"\xff\xd8\xff",
    ".png": b"\x89PNG\r\n\x1a\n",
    ".webp": b"RIFF",
}
NOMBRE_ADJUNTO = re.compile(r"^[0-9a-f]{32}\.(pdf|jpg|jpeg|png|webp)$")
# A partir de cuántos días una responsiva enviada a firmar se considera atrasada
DIAS_ESPERA_RESPONSIVA = 21
TAMANO_MAXIMO = 10 * 1024 * 1024  # 10 MB

PREFIJO_CODIGO = {
    TipoActivo.COMPUTO: "PC",
    TipoActivo.TELEFONO: "TEL",
    TipoActivo.TABLETA: "TAB",
    TipoActivo.LICENCIA_MICROSOFT: "MS",
    TipoActivo.LICENCIA_NORTON: "NOR",
    TipoActivo.CORREO_WORKSPACE: "GW",
}

# Sólo teléfonos, computadoras y tabletas generan responsiva
TIPOS_CON_RESPONSIVA = {TipoActivo.COMPUTO, TipoActivo.TABLETA, TipoActivo.TELEFONO}
# La INE sólo se pide para telefonía
TIPOS_CON_INE = {TipoActivo.TELEFONO}
# Activos son sólo los equipos físicos; licencias y correos son servicios asignados
TIPOS_FISICOS = [TipoActivo.COMPUTO, TipoActivo.TELEFONO, TipoActivo.TABLETA]
TIPOS_LICENCIA = {TipoActivo.LICENCIA_MICROSOFT, TipoActivo.LICENCIA_NORTON}
# Ningún dispositivo vence; las licencias toman la renovación de su familia
TIPOS_SIN_VENCIMIENTO = {
    TipoActivo.COMPUTO, TipoActivo.TELEFONO, TipoActivo.TABLETA, TipoActivo.CORREO_WORKSPACE
} | TIPOS_LICENCIA

OPERADOR_TELEFONIA = "Telcel"
ESQUEMA_TELEFONIA = "Plazo Forzoso con Equipo"
# Días sin iniciar sesión a partir de los cuales un correo se considera en desuso
DIAS_DESUSO_CORREO = 20

# Agrupaciones que usa la interfaz para separar el inventario por secciones
GRUPOS = {
    "equipos": [TipoActivo.COMPUTO],
    "tabletas": [TipoActivo.TABLETA],
    "telefonia": [TipoActivo.TELEFONO],
    "licencias": [TipoActivo.LICENCIA_MICROSOFT, TipoActivo.LICENCIA_NORTON],
    "correos": [TipoActivo.CORREO_WORKSPACE],
}


@asynccontextmanager
async def ciclo_vida(app: FastAPI):
    crear_esquema()
    _cifrar_adjuntos()
    with SessionLocal() as db:
        creados = sembrar(db)
        if any(creados.values()):
            print(f"[inventario] catálogos sembrados: {creados}")
    yield


app = FastAPI(
    title="Panel de control TI",
    version="0.3.0",
    lifespan=ciclo_vida,
    docs_url=None,
    redoc_url=None,
    openapi_url=None,
)


def ip_cliente(request: Request) -> str:
    return request.client.host if request.client else "desconocida"


def _cifrar_adjuntos() -> None:
    """Cifra los adjuntos que se guardaron antes de existir el cifrado."""
    for ruta in ADJUNTOS.iterdir():
        if ruta.is_file() and NOMBRE_ADJUNTO.match(ruta.name):
            contenido = ruta.read_bytes()
            if not es_cifrado(contenido):
                temporal = ruta.with_name(ruta.name + ".tmp")
                temporal.write_bytes(cifrar_bytes(contenido))
                temporal.replace(ruta)


CABECERAS_SEGURIDAD = {
    "X-Content-Type-Options": "nosniff",
    "X-Frame-Options": "DENY",
    "Referrer-Policy": "same-origin",
    "Permissions-Policy": "camera=(), microphone=(), geolocation=()",
    "Content-Security-Policy": (
        "default-src 'self'; img-src 'self' data:; style-src 'self' 'unsafe-inline'; "
        "script-src 'self' 'unsafe-inline'; object-src 'none'; base-uri 'self'; "
        "form-action 'self'; frame-ancestors 'none'"
    ),
}


@app.middleware("http")
async def revalidar_interfaz(request: Request, call_next):
    es_api = request.url.path.startswith("/api/")
    if es_api:
        espera = peticiones_por_ip.consumir(ip_cliente(request))
        if espera is not None:
            return JSONResponse(
                {"detail": "Demasiadas solicitudes. Intenta de nuevo en un momento."},
                status_code=429,
                headers={"Retry-After": str(int(espera))},
            )
    respuesta = await call_next(request)
    for nombre, valor in CABECERAS_SEGURIDAD.items():
        respuesta.headers.setdefault(nombre, valor)
    # Sin no-cache el navegador puede mezclar un index.html nuevo con CSS o JS viejos
    respuesta.headers["Cache-Control"] = "no-store" if es_api else "no-cache"
    return respuesta


def patron_busqueda(q: str) -> str:
    """Patrón LIKE en el que los comodines escritos por el usuario se buscan literalmente."""
    texto = q.strip().upper().replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
    return f"%{texto}%"


# --------------------------------------------------------------------------
# Utilidades
# --------------------------------------------------------------------------


# Los códigos consecutivos (EMP-0001, PC-0001) se calculan leyendo el último y sumando uno.
# Dos altas simultáneas leerían el mismo último, así que se hacen de una en una (una sola
# instancia atiende la aplicación, por eso basta un candado en memoria).
_candado_altas = threading.Lock()


def un_alta_a_la_vez(funcion):
    @functools.wraps(funcion)
    def envuelta(*args, **kwargs):
        with _candado_altas:
            return funcion(*args, **kwargs)

    return envuelta


def siguiente_codigo(db: Session, tipo: TipoActivo) -> str:
    prefijo = PREFIJO_CODIGO[tipo]
    ultimo = (
        db.query(Activo.codigo)
        .filter(Activo.codigo.like(f"{prefijo}-%"))
        .order_by(Activo.codigo.desc())
        .first()
    )
    consecutivo = 1
    if ultimo:
        try:
            consecutivo = int(ultimo[0].split("-")[-1]) + 1
        except ValueError:
            consecutivo = db.query(Activo).filter(Activo.tipo == tipo).count() + 1
    return f"{prefijo}-{consecutivo:04d}"


def siguiente_codigo_empleado(db: Session) -> str:
    ultimo = db.query(Empleado.codigo).order_by(Empleado.codigo.desc()).first()
    consecutivo = 1
    if ultimo:
        try:
            consecutivo = int(ultimo[0].split("-")[-1]) + 1
        except ValueError:
            consecutivo = db.query(Empleado).count() + 1
    return f"EMP-{consecutivo:04d}"


def asignacion_vigente(db: Session, activo_id: int) -> Asignacion | None:
    return (
        db.query(Asignacion)
        .filter(Asignacion.activo_id == activo_id, Asignacion.fecha_devolucion.is_(None))
        .first()
    )


def correos_asignados(db: Session, empleado_ids: list[int]) -> dict[int, list[str]]:
    """Cuentas de Workspace vigentes de cada persona: el correo sale de lo asignado."""
    if not empleado_ids:
        return {}
    salida: dict[int, list[str]] = {}
    for empleado_id, direccion in (
        db.query(Asignacion.empleado_id, DetalleCorreo.direccion)
        .join(DetalleCorreo, DetalleCorreo.activo_id == Asignacion.activo_id)
        .filter(Asignacion.empleado_id.in_(empleado_ids), Asignacion.fecha_devolucion.is_(None))
        .order_by(DetalleCorreo.direccion)
    ):
        salida.setdefault(empleado_id, []).append(direccion)
    return salida


def filtro_ubicacion(
    sucursal_id: int | None,
    departamento_id: int | None,
    bodega_id: int | None,
    planta_id: int | None,
) -> list:
    """Condiciones sobre Empleado para cada catálogo de ubicación elegido."""
    condiciones = []
    if sucursal_id:
        condiciones.append(Empleado.sucursal_id == sucursal_id)
    if departamento_id:
        condiciones.append(Empleado.departamento_id == departamento_id)
    if bodega_id:
        condiciones.append(Empleado.bodega_id == bodega_id)
    if planta_id:
        condiciones.append(Empleado.planta_id == planta_id)
    return condiciones


def dic_empleado(e: Empleado, con_activos: bool = False, db: Session | None = None) -> dict:
    datos = {
        "id": e.id,
        "codigo": e.codigo,
        "nombres": e.nombres,
        "apellidos": e.apellidos,
        "nombre_completo": e.nombre_completo,
        "puesto": e.puesto,
        "sucursal_id": e.sucursal_id,
        "sucursal": e.sucursal.nombre if e.sucursal else None,
        "departamento_id": e.departamento_id,
        "departamento": e.departamento.nombre if e.departamento else None,
        "bodega_id": e.bodega_id,
        "bodega": e.bodega.nombre if e.bodega else None,
        "planta_id": e.planta_id,
        "planta": e.planta.nombre if e.planta else None,
        "estatus": e.estatus.value,
        "fecha_ingreso": e.fecha_ingreso,
        "fecha_baja": e.fecha_baja,
        "notas": e.notas,
    }
    if con_activos and db is not None:
        vigentes = (
            db.query(Asignacion)
            .filter(Asignacion.empleado_id == e.id, Asignacion.fecha_devolucion.is_(None))
            .all()
        )
        datos["activos"] = [dic_asignacion(a) for a in vigentes]
        datos["correos"] = correos_asignados(db, [e.id]).get(e.id, [])
        historial = (
            db.query(Asignacion)
            .filter(Asignacion.empleado_id == e.id, Asignacion.fecha_devolucion.isnot(None))
            .order_by(Asignacion.fecha_devolucion.desc())
            .all()
        )
        datos["historial"] = [dic_asignacion(a) for a in historial]
        datos["accesos"] = accesos_de(db, e.id)
    return datos


def accesos_de(db: Session, empleado_id: int) -> list[dict]:
    """Checklist de sistemas internos, con todos los sistemas activos presentes."""
    registrados = {
        a.sistema_id: a
        for a in db.query(AccesoSistema).filter(AccesoSistema.empleado_id == empleado_id)
    }
    salida = []
    for s in (
        db.query(SistemaInterno)
        .filter_by(activo=True)
        .order_by(SistemaInterno.orden, SistemaInterno.nombre)
    ):
        a = registrados.get(s.id)
        salida.append(
            {
                "sistema_id": s.id,
                "sistema": s.nombre,
                "clave": s.clave,
                "descripcion": s.descripcion,
                "entregado": bool(a and a.entregado),
                "fecha_entrega": a.fecha_entrega if a else None,
                "usuario_sistema": a.usuario_sistema if a else None,
                "notas": a.notas if a else None,
                "actualizado_por": a.actualizado_por if a else None,
                "actualizado_en": a.actualizado_en if a else None,
            }
        )
    return salida


def dic_activo(
    a: Activo, db: Session | None = None, vigentes: dict[int, Asignacion] | None = None
) -> dict:
    datos = {
        "id": a.id,
        "codigo": a.codigo,
        "tipo": a.tipo.value,
        "estatus": a.estatus.value,
        "descripcion": a.descripcion,
        "fecha_adquisicion": a.fecha_adquisicion,
        "fecha_vencimiento": a.fecha_vencimiento,
        "lleva_responsiva": a.tipo in TIPOS_CON_RESPONSIVA,
        "lleva_ine": a.tipo in TIPOS_CON_INE,
        "notas": a.notas,
    }
    if a.detalle_computo:
        d = a.detalle_computo
        datos["computo"] = {
            "marca": d.marca,
            "modelo": d.modelo,
            "numero_serie": d.numero_serie,
            "procesador": d.procesador,
            "ram_gb": d.ram_gb,
            "almacenamiento": d.almacenamiento,
            "sistema_operativo": d.sistema_operativo,
            "clave_windows": d.clave_windows,
        }
    if a.detalle_movil:
        d = a.detalle_movil
        datos["movil"] = {
            "marca": d.marca,
            "modelo": d.modelo,
            "imei": d.imei,
            "numero_telefonico": d.numero_telefonico,
            "iccid": d.iccid,
            "operador": d.operador,
            "precio_plan": float(d.precio_plan) if d.precio_plan is not None else None,
            "esquema": d.esquema,
            "esim": d.esim,
            "cuenta_operador": d.cuenta_operador,
        }
    if a.detalle_licencia:
        d = a.detalle_licencia
        datos["licencia"] = {
            "cuenta_id": d.cuenta_id,
            "familia": d.cuenta.nombre if d.cuenta else None,
            "correo_familia": d.cuenta.correo_cuenta if d.cuenta else None,
            "ranuras_familia": d.cuenta.ranuras_totales if d.cuenta else None,
            "correo_licencia": d.correo_licencia,
            "identificador_dispositivo": d.identificador_dispositivo,
        }
    if a.detalle_correo:
        d = a.detalle_correo
        datos["correo"] = {
            "direccion": d.direccion,
            "habilitada": d.habilitada,
            "ultimo_acceso": d.ultimo_acceso,
            "almacenamiento_gb": float(d.almacenamiento_gb)
            if d.almacenamiento_gb is not None
            else None,
            "alias_de": d.alias_de,
            "actividad_importada_en": d.actividad_importada_en,
            "dias_sin_acceso": dias_sin_acceso(d),
            "nunca_accedio": d.actividad_importada_en is not None and d.ultimo_acceso is None,
        }
    if vigentes is not None:
        vig = vigentes.get(a.id)
        datos["asignado_a"] = dic_asignacion(vig) if vig else None
    elif db is not None:
        vig = asignacion_vigente(db, a.id)
        datos["asignado_a"] = dic_asignacion(vig) if vig else None
    return datos


def asignaciones_vigentes(db: Session, activo_ids: list[int]) -> dict[int, Asignacion]:
    """Asignación vigente de cada activo en pocas consultas, en lugar de una por activo."""
    vigentes: dict[int, Asignacion] = {}
    for i in range(0, len(activo_ids), 500):
        lote = activo_ids[i : i + 500]
        for asig in (
            db.query(Asignacion)
            .options(joinedload(Asignacion.empleado))
            .filter(Asignacion.activo_id.in_(lote), Asignacion.fecha_devolucion.is_(None))
        ):
            vigentes.setdefault(asig.activo_id, asig)
    return vigentes


def dias_sin_acceso(d: DetalleCorreo) -> int | None:
    if d.ultimo_acceso is None:
        return None
    return (ahora_local() - d.ultimo_acceso).days


def correo_en_desuso(d: DetalleCorreo) -> bool:
    if not d.habilitada:
        return False
    dias = dias_sin_acceso(d)
    if dias is None:
        return d.actividad_importada_en is not None
    return dias >= DIAS_DESUSO_CORREO


def dic_asignacion(a: Asignacion) -> dict:
    return {
        "id": a.id,
        "activo_id": a.activo_id,
        "activo_codigo": a.activo.codigo,
        "activo_tipo": a.activo.tipo.value,
        "activo_descripcion": a.activo.descripcion,
        "empleado_id": a.empleado_id,
        "empleado": a.empleado.nombre_completo,
        "codigo_empleado": a.empleado.codigo,
        "fecha_asignacion": a.fecha_asignacion,
        "fecha_devolucion": a.fecha_devolucion,
        "responsiva_estado": a.responsiva_estado.value,
        "responsiva_fecha_envio": a.responsiva_fecha_envio,
        "responsiva_fecha_firma": a.responsiva_fecha_firma,
        "responsiva_destino": a.responsiva_destino,
        "responsiva_url": a.responsiva_url,
        "responsiva_archivo": a.responsiva_archivo,
        "responsiva_nombre": a.responsiva_nombre,
        "dias_en_espera": (
            (hoy_local() - a.responsiva_fecha_envio).days
            if a.responsiva_estado == EstadoDocumento.ENVIADA and a.responsiva_fecha_envio
            else None
        ),
        "ine_estado": a.ine_estado.value,
        "ine_url": a.ine_url,
        "ine_archivo": a.ine_archivo,
        "ine_nombre": a.ine_nombre,
        "motivo_devolucion": a.motivo_devolucion,
        "notas": a.notas,
    }


# --------------------------------------------------------------------------
# Sesion
# --------------------------------------------------------------------------


def _perfil(usuario: Usuario) -> dict:
    return {
        "id": usuario.id,
        "username": usuario.username,
        "nombre": usuario.nombre,
        "rol": usuario.rol.value,
        "foto": usuario.foto,
        "correo": usuario.correo,
        "telefono": usuario.telefono,
        "puesto": usuario.puesto,
        "password_por_defecto": usuario.debe_cambiar_password,
    }


@app.post("/api/login")
def login(
    datos: schemas.LoginIn,
    request: Request,
    respuesta: Response,
    db: Session = Depends(obtener_sesion),
):
    username = datos.username.strip().lower()
    ip = ip_cliente(request)
    llave_usuario = f"{ip}|{username}"
    espera = fallos_por_ip.excedido(ip) or fallos_por_usuario.excedido(llave_usuario)
    if espera is not None:
        raise HTTPException(
            status.HTTP_429_TOO_MANY_REQUESTS,
            f"Demasiados intentos fallidos. Intenta de nuevo en {int(espera // 60) + 1} min.",
            headers={"Retry-After": str(int(espera))},
        )

    usuario = db.query(Usuario).filter(Usuario.username == username).first()
    if usuario is None:
        verificar_senuelo(datos.password)
    if usuario is None or not verificar_password(datos.password, usuario.password_hash):
        fallos_por_ip.registrar(ip)
        fallos_por_usuario.registrar(llave_usuario)
        registrar(
            db, None, "LOGIN_FALLIDO", "usuario", usuario.id if usuario else None, f"{username} · {ip}"
        )
        db.commit()
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Usuario o contraseña incorrectos")
    if not usuario.activo:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Usuario desactivado")

    fallos_por_usuario.limpiar(llave_usuario)
    if requiere_rehash(usuario.password_hash):
        usuario.password_hash = hashear_password(datos.password)
    # Bases anteriores: si conserva la contraseña de fábrica, debe cambiarla
    if not usuario.debe_cambiar_password and datos.password == PASSWORD_DE_FABRICA:
        usuario.debe_cambiar_password = True
    token = crear_sesion(db, usuario)
    poner_cookie(respuesta, token)
    return _perfil(usuario)


@app.post("/api/logout")
def logout(request: Request, respuesta: Response, db: Session = Depends(obtener_sesion)):
    token = request.cookies.get(COOKIE)
    if token:
        cerrar_sesion(db, token)
    quitar_cookie(respuesta)
    return {"ok": True}


@app.get("/api/yo")
def yo(usuario: Usuario = Depends(usuario_actual)):
    return _perfil(usuario)


@app.put("/api/yo")
def editar_perfil(
    datos: schemas.PerfilIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(usuario_actual),
):
    usuario.nombre = datos.nombre.strip()
    usuario.correo = datos.correo
    usuario.telefono = datos.telefono
    usuario.puesto = datos.puesto
    registrar(db, usuario, "EDITAR", "usuario", usuario.id, "perfil")
    db.commit()
    return _perfil(usuario)


FOTO_MAXIMA = 2 * 1024 * 1024  # 2 MB


def _quitar_archivo_foto(nombre: str | None) -> None:
    if nombre and NOMBRE_ADJUNTO.match(nombre):
        (ADJUNTOS / nombre).unlink(missing_ok=True)


@app.post("/api/yo/foto")
async def subir_foto(
    archivo: UploadFile = File(...),
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(usuario_actual),
):
    extension = Path(archivo.filename or "").suffix.lower()
    if extension not in {".jpg", ".jpeg", ".png", ".webp"}:
        raise HTTPException(400, "La foto debe ser JPG, PNG o WEBP")
    contenido = await archivo.read(FOTO_MAXIMA + 1)
    if len(contenido) > FOTO_MAXIMA:
        raise HTTPException(400, "La foto supera los 2 MB")
    if not contenido.startswith(FIRMAS_ARCHIVO[extension]) or (
        extension == ".webp" and contenido[8:12] != b"WEBP"
    ):
        raise HTTPException(400, "El contenido del archivo no corresponde a su formato")

    nombre_disco = f"{secrets.token_hex(16)}{extension}"
    (ADJUNTOS / nombre_disco).write_bytes(cifrar_bytes(contenido))
    _quitar_archivo_foto(usuario.foto)
    usuario.foto = nombre_disco
    registrar(db, usuario, "EDITAR", "usuario", usuario.id, "foto de perfil")
    db.commit()
    return _perfil(usuario)


@app.delete("/api/yo/foto")
def quitar_foto(
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(usuario_actual),
):
    _quitar_archivo_foto(usuario.foto)
    usuario.foto = None
    registrar(db, usuario, "EDITAR", "usuario", usuario.id, "se quitó la foto de perfil")
    db.commit()
    return _perfil(usuario)


@app.post("/api/yo/password")
def cambiar_password(
    datos: schemas.CambioPasswordIn,
    request: Request,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(usuario_actual),
):
    llave = f"password|{usuario.id}"
    if fallos_por_usuario.excedido(llave) is not None:
        raise HTTPException(status.HTTP_429_TOO_MANY_REQUESTS, "Demasiados intentos. Intenta más tarde.")
    if not verificar_password(datos.password_actual, usuario.password_hash):
        fallos_por_usuario.registrar(llave)
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "La contraseña actual no coincide")
    if datos.password_nueva == PASSWORD_DE_FABRICA:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "La contraseña nueva no puede ser la de fábrica")
    usuario.password_hash = hashear_password(datos.password_nueva)
    usuario.debe_cambiar_password = False
    cerrar_otras_sesiones(db, usuario, request.cookies.get(COOKIE))
    registrar(db, usuario, "EDITAR", "usuario", usuario.id, "cambio de contraseña")
    db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------
# Adjuntos: responsivas e INE
# --------------------------------------------------------------------------


@app.post("/api/adjuntos")
async def subir_adjunto(
    request: Request,
    archivo: UploadFile = File(...),
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    if subidas_por_ip.consumir(ip_cliente(request)) is not None:
        raise HTTPException(429, "Demasiados archivos en poco tiempo. Intenta en un momento.")
    extension = Path(archivo.filename or "").suffix.lower()
    if extension not in EXTENSIONES_PERMITIDAS:
        raise HTTPException(
            400, f"Formato no admitido. Se aceptan: {', '.join(sorted(EXTENSIONES_PERMITIDAS))}"
        )

    contenido = await archivo.read(TAMANO_MAXIMO + 1)
    if len(contenido) > TAMANO_MAXIMO:
        raise HTTPException(400, "El archivo supera los 10 MB")
    if not contenido:
        raise HTTPException(400, "El archivo está vacío")
    if not contenido.startswith(FIRMAS_ARCHIVO[extension]) or (
        extension == ".webp" and contenido[8:12] != b"WEBP"
    ):
        raise HTTPException(400, "El contenido del archivo no corresponde a su formato")
    nombre_original = re.sub(r"[^\w .()-]", "_", Path(archivo.filename or "").name)[:200]

    # El nombre en disco lo genera el servidor: nunca se usa el del usuario,
    # que podría contener rutas o caracteres peligrosos.
    nombre_disco = f"{secrets.token_hex(16)}{extension}"
    (ADJUNTOS / nombre_disco).write_bytes(cifrar_bytes(contenido))

    registrar(db, usuario, "SUBIR", "adjunto", None, nombre_original)
    db.commit()
    return {
        "archivo": nombre_disco,
        "nombre": nombre_original,
        "tamano": len(contenido),
    }


@app.get("/api/adjuntos/{nombre}")
def ver_adjunto(nombre: str, _=Depends(usuario_actual)):
    # Solo nombres generados por el servidor: impide rutas tipo ../../
    ruta = ADJUNTOS / nombre
    if not NOMBRE_ADJUNTO.match(nombre) or not ruta.is_file():
        raise HTTPException(404, "Archivo no encontrado")

    tipo, _codificacion = mimetypes.guess_type(ruta.name)
    return Response(
        descifrar_bytes(ruta.read_bytes()),
        media_type=tipo or "application/octet-stream",
        # inline permite verlo dentro de la aplicación en lugar de descargarlo
        headers={"Content-Disposition": f'inline; filename="{ruta.name}"'},
    )


# --------------------------------------------------------------------------
# Catalogos
# --------------------------------------------------------------------------


@app.get("/api/catalogos")
def catalogos(db: Session = Depends(obtener_sesion), _=Depends(usuario_actual)):
    return {
        "sucursales": [
            {"id": s.id, "nombre": s.nombre, "clave": s.clave}
            for s in db.query(Sucursal).filter_by(activo=True).order_by(Sucursal.nombre)
        ],
        "departamentos": [
            {"id": d.id, "nombre": d.nombre, "clave": d.clave}
            for d in db.query(Departamento).filter_by(activo=True).order_by(Departamento.nombre)
        ],
        "bodegas": [
            {"id": b.id, "nombre": b.nombre, "sucursal_id": b.sucursal_id}
            for b in db.query(Bodega).filter_by(activo=True).order_by(Bodega.nombre)
        ],
        "plantas": [
            {"id": p.id, "nombre": p.nombre}
            for p in db.query(Planta).filter_by(activo=True).order_by(Planta.nombre)
        ],
        "cuentas_licencia": [
            {
                "id": c.id,
                "nombre": c.nombre,
                "producto": c.producto,
                "correo_cuenta": c.correo_cuenta,
                "referencia_credencial": c.referencia_credencial,
                "ranuras_totales": c.ranuras_totales,
                "ranuras_usadas": db.query(DetalleLicencia).filter_by(cuenta_id=c.id).count(),
                "fecha_renovacion": c.fecha_renovacion,
                "notas": c.notas,
            }
            for c in db.query(CuentaLicencia).filter_by(activo=True).order_by(CuentaLicencia.nombre)
        ],
        "sistemas": [
            {
                "id": s.id,
                "nombre": s.nombre,
                "clave": s.clave,
                "descripcion": s.descripcion,
                "orden": s.orden,
            }
            for s in db.query(SistemaInterno)
            .filter_by(activo=True)
            .order_by(SistemaInterno.orden, SistemaInterno.nombre)
        ],
        "analistas": [
            {"id": u.id, "nombre": u.nombre, "username": u.username, "rol": u.rol.value}
            for u in db.query(Usuario)
            .filter(Usuario.activo.is_(True))
            .order_by(Usuario.nombre)
        ],
        "tipos_activo": [t.value for t in TipoActivo],
        "estatus_activo": [e.value for e in EstatusActivo],
    }


@app.post("/api/catalogos/{entidad}")
def crear_catalogo(
    entidad: str,
    datos: schemas.CatalogoIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    modelos = {
        "sucursales": Sucursal,
        "departamentos": Departamento,
        "bodegas": Bodega,
        "plantas": Planta,
    }
    if entidad not in modelos:
        raise HTTPException(404, "Catálogo desconocido")

    Modelo = modelos[entidad]
    kwargs: dict = {"nombre": datos.nombre.strip().upper()}
    if entidad in ("sucursales", "departamentos"):
        if not datos.clave:
            raise HTTPException(400, "La clave es obligatoria")
        kwargs["clave"] = datos.clave.strip().upper()
    if entidad == "bodegas":
        kwargs["sucursal_id"] = datos.sucursal_id

    registro = Modelo(**kwargs)
    db.add(registro)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Ya existe un registro con ese nombre o clave")
    registrar(db, usuario, "CREAR", entidad, registro.id, datos.nombre)
    db.commit()
    return {"id": registro.id, "nombre": registro.nombre}


# Tope de licencias por familia según el plan de cada producto
LICENCIAS_POR_FAMILIA = {TipoActivo.LICENCIA_MICROSOFT: 6, TipoActivo.LICENCIA_NORTON: 10}
NOMBRE_PRODUCTO = {TipoActivo.LICENCIA_MICROSOFT: "Microsoft 365", TipoActivo.LICENCIA_NORTON: "Norton 360"}


def tipo_de_producto(producto: str | None) -> TipoActivo | None:
    texto = (producto or "").lower()
    if "norton" in texto:
        return TipoActivo.LICENCIA_NORTON
    if "microsoft" in texto or "365" in texto or "office" in texto:
        return TipoActivo.LICENCIA_MICROSOFT
    return None


def validar_ranuras(datos: schemas.CuentaLicenciaIn, en_uso: int = 0) -> None:
    tipo = tipo_de_producto(datos.producto)
    if tipo and datos.ranuras_totales > LICENCIAS_POR_FAMILIA[tipo]:
        raise HTTPException(
            400,
            f"Una familia de {NOMBRE_PRODUCTO[tipo]} admite hasta "
            f"{LICENCIAS_POR_FAMILIA[tipo]} licencias",
        )
    if datos.ranuras_totales < en_uso:
        raise HTTPException(
            409, f"La familia ya tiene {en_uso} licencias; no puede quedar con menos ranuras"
        )


def validar_cupo_familia(db: Session, activo: Activo, tipo: TipoActivo, cuenta_id: int) -> None:
    """Impide ligar una licencia a una familia llena o de otro producto."""
    cuenta = db.get(CuentaLicencia, cuenta_id)
    if cuenta is None:
        raise HTTPException(404, "Familia de licencias no encontrada")
    producto = tipo_de_producto(cuenta.producto)
    if producto and producto != tipo:
        raise HTTPException(
            409, f"{cuenta.nombre} es una familia de {NOMBRE_PRODUCTO[producto]}"
        )
    tope = min(LICENCIAS_POR_FAMILIA[tipo], cuenta.ranuras_totales)
    ocupadas = (
        db.query(DetalleLicencia)
        .filter(DetalleLicencia.cuenta_id == cuenta_id, DetalleLicencia.activo_id != activo.id)
        .count()
    )
    if ocupadas >= tope:
        raise HTTPException(
            409, f"{cuenta.nombre} ya tiene sus {tope} licencias; elige otra familia"
        )


@app.post("/api/cuentas-licencia")
def crear_cuenta_licencia(
    datos: schemas.CuentaLicenciaIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    validar_ranuras(datos)
    cuenta = CuentaLicencia(**datos.model_dump())
    db.add(cuenta)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Ya existe una cuenta con ese nombre")
    registrar(db, usuario, "CREAR", "cuenta_licencia", cuenta.id, datos.nombre)
    db.commit()
    return {"id": cuenta.id, "nombre": cuenta.nombre}


@app.put("/api/cuentas-licencia/{cuenta_id}")
def editar_cuenta_licencia(
    cuenta_id: int,
    datos: schemas.CuentaLicenciaIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    cuenta = db.get(CuentaLicencia, cuenta_id)
    if cuenta is None:
        raise HTTPException(404, "Cuenta no encontrada")
    validar_ranuras(datos, db.query(DetalleLicencia).filter_by(cuenta_id=cuenta_id).count())
    for campo, valor in datos.model_dump().items():
        setattr(cuenta, campo, valor)
    registrar(db, usuario, "EDITAR", "cuenta_licencia", cuenta.id, cuenta.nombre)
    db.commit()
    return {"id": cuenta.id, "nombre": cuenta.nombre}


@app.delete("/api/cuentas-licencia/{cuenta_id}")
def eliminar_cuenta_licencia(
    cuenta_id: int,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    cuenta = db.get(CuentaLicencia, cuenta_id)
    if cuenta is None:
        raise HTTPException(404, "Cuenta no encontrada")
    en_uso = db.query(DetalleLicencia).filter_by(cuenta_id=cuenta_id).count()
    if en_uso:
        raise HTTPException(
            409, f"La cuenta tiene {en_uso} licencia(s) registrada(s). Elimínalas primero."
        )
    nombre = cuenta.nombre
    db.delete(cuenta)
    registrar(db, usuario, "ELIMINAR", "cuenta_licencia", cuenta_id, nombre)
    db.commit()
    return {"ok": True}


PRODUCTOS_LICENCIA = {
    "microsoft": (TipoActivo.LICENCIA_MICROSOFT, "microsoft"),
    "norton": (TipoActivo.LICENCIA_NORTON, "norton"),
}


@app.get("/api/licencias/familias")
def licencias_por_familia(
    producto: str = Query("microsoft", pattern="^(microsoft|norton)$"),
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    """Licencias agrupadas por familia (cuenta), con quién las usa y dónde está."""
    tipo, palabra = PRODUCTOS_LICENCIA[producto]
    licencias = (
        db.query(Activo)
        .join(DetalleLicencia)
        .filter(Activo.tipo == tipo)
        .order_by(Activo.codigo)
        .all()
    )
    grupos: dict[int | None, list[dict]] = {}
    for a in licencias:
        d = a.detalle_licencia
        vig = asignacion_vigente(db, a.id)
        e = vig.empleado if vig else None
        grupos.setdefault(d.cuenta_id, []).append(
            {
                "id": a.id,
                "codigo": a.codigo,
                "estatus": a.estatus.value,
                "correo_licencia": d.correo_licencia,
                "identificador": d.identificador_dispositivo,
                "fecha_adquisicion": a.fecha_adquisicion,
                "fecha_vencimiento": a.fecha_vencimiento,
                "asignacion_id": vig.id if vig else None,
                "empleado_id": e.id if e else None,
                "empleado": e.nombre_completo if e else None,
                "sucursal": e.sucursal.nombre if e and e.sucursal else None,
                "bodega": e.bodega.nombre if e and e.bodega else None,
                "departamento": e.departamento.nombre if e and e.departamento else None,
                "planta": e.planta.nombre if e and e.planta else None,
                "sucursal_id": e.sucursal_id if e else None,
                "departamento_id": e.departamento_id if e else None,
                "bodega_id": e.bodega_id if e else None,
                "planta_id": e.planta_id if e else None,
            }
        )

    cuentas = [
        c
        for c in db.query(CuentaLicencia).order_by(CuentaLicencia.id)
        if c.id in grupos or (c.activo and palabra in (c.producto or "").lower())
    ]
    salida = [
        {
            "cuenta": {
                "id": c.id,
                "nombre": c.nombre,
                "producto": c.producto,
                "correo_cuenta": c.correo_cuenta,
                "referencia_credencial": c.referencia_credencial,
                "ranuras_totales": c.ranuras_totales,
                "fecha_renovacion": c.fecha_renovacion,
                "notas": c.notas,
            },
            "licencias": grupos.get(c.id, []),
        }
        for c in cuentas
    ]
    if grupos.get(None):
        salida.append({"cuenta": None, "licencias": grupos[None]})
    return salida


# --------------------------------------------------------------------------
# Empleados
# --------------------------------------------------------------------------


@app.get("/api/empleados")
def listar_empleados(
    q: str | None = Query(None, max_length=100),
    sucursal_id: int | None = None,
    departamento_id: int | None = None,
    bodega_id: int | None = None,
    planta_id: int | None = None,
    estatus: EstatusEmpleado | None = None,
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    consulta = db.query(Empleado)
    if q:
        patron = patron_busqueda(q)
        consulta = consulta.filter(
            or_(
                func.upper(Empleado.nombre_completo).like(patron, escape="\\"),
                func.upper(Empleado.codigo).like(patron, escape="\\"),
                Empleado.id.in_(
                    db.query(Asignacion.empleado_id)
                    .join(DetalleCorreo, DetalleCorreo.activo_id == Asignacion.activo_id)
                    .filter(
                        Asignacion.fecha_devolucion.is_(None),
                        func.upper(DetalleCorreo.direccion).like(patron, escape="\\"),
                    )
                ),
            )
        )
    consulta = consulta.filter(*filtro_ubicacion(sucursal_id, departamento_id, bodega_id, planta_id))
    if estatus:
        consulta = consulta.filter(Empleado.estatus == estatus)

    empleados = consulta.order_by(Empleado.nombre_completo).all()
    ids = [e.id for e in empleados]
    def contar(fisicos: bool) -> dict:
        if not ids:
            return {}
        filtro = Activo.tipo.in_(TIPOS_FISICOS) if fisicos else Activo.tipo.notin_(TIPOS_FISICOS)
        return dict(
            db.query(Asignacion.empleado_id, func.count(Asignacion.id))
            .join(Activo, Activo.id == Asignacion.activo_id)
            .filter(
                Asignacion.empleado_id.in_(ids),
                Asignacion.fecha_devolucion.is_(None),
                filtro,
            )
            .group_by(Asignacion.empleado_id)
            .all()
        )

    activos, servicios = contar(True), contar(False)
    correos = correos_asignados(db, ids)
    salida = []
    for e in empleados:
        d = dic_empleado(e)
        d["total_activos"] = activos.get(e.id, 0)
        d["total_servicios"] = servicios.get(e.id, 0)
        d["correos"] = correos.get(e.id, [])
        salida.append(d)
    return salida


ESTADO_DOC_TEXTO = {
    "PENDIENTE": "Pendiente de enviar",
    "ENVIADA": "En espera de firma",
    "FIRMADA": "Firmada",
    "NO_APLICA": "No aplica",
}

TIPO_ACTIVO_TEXTO = {
    TipoActivo.COMPUTO: "Equipo de cómputo",
    TipoActivo.TABLETA: "Tablet",
    TipoActivo.TELEFONO: "Teléfono",
}


def _neutralizar_formulas(wb) -> None:
    """Un texto que empieza con = + - @ se abriría como fórmula en Excel: se antepone una comilla."""
    for hoja in wb.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.data_type in ("s", "f") and isinstance(celda.value, str) \
                        and celda.value[:1] in ("=", "+", "-", "@", "\t", "\r"):
                    celda.value = "'" + celda.value
                    celda.data_type = "s"


def _descarga_excel(hoja_titulo: str, columnas: list[str], filas: list[list], anchos: list[int], nombre: str):
    """Una hoja con encabezado fijo y filtros. Una celda (texto, url) se escribe como enlace."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    hoja = wb.active
    hoja.title = hoja_titulo
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="DDEBF7")
        celda.alignment = Alignment(vertical="center", wrap_text=True)

    for n, fila in enumerate(filas, start=2):
        for col, valor in enumerate(fila, start=1):
            if isinstance(valor, tuple):
                texto, url = valor
                celda = hoja.cell(row=n, column=col, value=texto)
                if url:
                    celda.hyperlink = url
                    celda.font = Font(color="0563C1", underline="single")
            else:
                celda = hoja.cell(row=n, column=col, value=valor)
            if isinstance(valor, dt.date):
                celda.number_format = "DD/MM/YYYY"
            celda.alignment = Alignment(vertical="top", wrap_text=True)

    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    _neutralizar_formulas(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}_{hoy_local():%Y%m%d}.xlsx"'},
    )


def _enlace_documento(request: Request, asignacion: dict | None, cual: str):
    """Liga del PDF o imagen de una responsiva o INE, o vacío si no hay."""
    if not asignacion:
        return ""
    archivo = asignacion[f"{cual}_archivo"]
    if archivo:
        url = f"{str(request.base_url).rstrip('/')}/api/adjuntos/{quote(archivo)}"
        return (asignacion[f"{cual}_nombre"] or archivo, url)
    url = asignacion[f"{cual}_url"]
    return (url, url) if url else ""


@app.get("/api/empleados/exportar")
def exportar_empleados(
    q: str | None = Query(None, max_length=100),
    sucursal_id: int | None = None,
    departamento_id: int | None = None,
    bodega_id: int | None = None,
    planta_id: int | None = None,
    estatus: EstatusEmpleado | None = None,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(usuario_actual),
):
    empleados = listar_empleados(
        q, sucursal_id, departamento_id, bodega_id, planta_id, estatus, db, usuario
    )
    ids = [e["id"] for e in empleados]
    equipos: dict[int, list[str]] = {}
    if ids:
        for a in (
            db.query(Asignacion)
            .join(Activo, Activo.id == Asignacion.activo_id)
            .filter(
                Asignacion.empleado_id.in_(ids),
                Asignacion.fecha_devolucion.is_(None),
                Activo.tipo.in_(TIPOS_FISICOS),
            )
            .order_by(Activo.codigo)
        ):
            texto = f"{a.activo.codigo} ({TIPO_ACTIVO_TEXTO.get(a.activo.tipo, a.activo.tipo.value)})"
            if a.activo.tipo in TIPOS_CON_RESPONSIVA:
                texto += f" · Responsiva: {ESTADO_DOC_TEXTO[a.responsiva_estado.value]}"
            equipos.setdefault(a.empleado_id, []).append(texto)

    filas = [
        [
            e["codigo"], e["nombre_completo"], "\n".join(e["correos"]), e["puesto"],
            e["departamento"], e["sucursal"], e["bodega"] or e["planta"], e["estatus"],
            e["fecha_ingreso"], e["fecha_baja"], e["total_activos"],
            "\n".join(equipos.get(e["id"], [])), e["total_servicios"],
        ]
        for e in empleados
    ]
    return _descarga_excel(
        "Empleados",
        ["Código", "Nombre completo", "Correos", "Puesto", "Departamento", "Sucursal",
         "Locación", "Estatus", "Fecha de ingreso", "Fecha de baja", "Activos",
         "Equipos asignados", "Licencias y correos"],
        filas,
        [12, 34, 32, 24, 22, 18, 24, 11, 14, 14, 9, 52, 12],
        "empleados",
    )


@app.get("/api/empleados/{empleado_id}")
def ver_empleado(
    empleado_id: int, db: Session = Depends(obtener_sesion), _=Depends(usuario_actual)
):
    empleado = db.get(Empleado, empleado_id)
    if empleado is None:
        raise HTTPException(404, "Empleado no encontrado")
    return dic_empleado(empleado, con_activos=True, db=db)


def lugar_trabajo(e: Empleado) -> str:
    if e.bodega:
        return e.bodega.nombre
    return e.planta.nombre if e.planta else ""


def nombre_completo(datos: schemas.EmpleadoIn) -> str:
    return " ".join(f"{datos.apellidos} {datos.nombres}".split()).upper()


def validar_lugar_trabajo(datos: schemas.EmpleadoIn) -> None:
    if datos.bodega_id and datos.planta_id:
        raise HTTPException(422, "Un empleado solo puede tener bodega o planta, no ambas")


def validar_empleado_unico(db: Session, clave: str, excluir_id: int | None = None) -> None:
    consulta = db.query(Empleado).filter(Empleado.nombre_clave == clave)
    if excluir_id is not None:
        consulta = consulta.filter(Empleado.id != excluir_id)
    existente = consulta.first()
    if existente:
        raise HTTPException(
            409,
            f"{existente.nombre_completo} ya está registrado como {existente.codigo} "
            f"({existente.estatus.value})",
        )


@app.post("/api/empleados")
@un_alta_a_la_vez
def crear_empleado(
    datos: schemas.EmpleadoIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    validar_lugar_trabajo(datos)
    clave = clave_nombre(datos.nombres, datos.apellidos)
    validar_empleado_unico(db, clave)
    empleado = Empleado(
        codigo=siguiente_codigo_empleado(db),
        nombre_completo=nombre_completo(datos),
        nombre_clave=clave,
        **datos.model_dump(),
    )
    db.add(empleado)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Ya existe un empleado con ese código o nombre")
    registrar(db, usuario, "CREAR", "empleado", empleado.id, empleado.nombre_completo)
    db.commit()
    return dic_empleado(empleado)


@app.put("/api/empleados/{empleado_id}")
def editar_empleado(
    empleado_id: int,
    datos: schemas.EmpleadoIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    empleado = db.get(Empleado, empleado_id)
    if empleado is None:
        raise HTTPException(404, "Empleado no encontrado")
    validar_lugar_trabajo(datos)
    clave = clave_nombre(datos.nombres, datos.apellidos)
    validar_empleado_unico(db, clave, excluir_id=empleado_id)
    for campo, valor in datos.model_dump().items():
        setattr(empleado, campo, valor)
    empleado.nombre_completo = nombre_completo(datos)
    empleado.nombre_clave = clave
    registrar(db, usuario, "EDITAR", "empleado", empleado.id, empleado.nombre_completo)
    db.commit()
    return dic_empleado(empleado)


@app.delete("/api/empleados/{empleado_id}")
def eliminar_empleado(
    empleado_id: int,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    empleado = db.get(Empleado, empleado_id)
    if empleado is None:
        raise HTTPException(404, "Empleado no encontrado")

    total = db.query(Asignacion).filter(Asignacion.empleado_id == empleado_id).count()
    if total:
        raise HTTPException(
            409,
            f"Tiene {total} asignación(es) en su historial. "
            "Para conservar el rastro de los activos, márcalo como BAJA en lugar de eliminarlo.",
        )
    nombre = empleado.nombre_completo
    db.delete(empleado)
    registrar(db, usuario, "ELIMINAR", "empleado", empleado_id, nombre)
    db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------
# Activos
# --------------------------------------------------------------------------


def _aplicar_detalle(db: Session, activo: Activo, datos: schemas.ActivoIn) -> None:
    """Crea o actualiza las tablas de detalle que corresponden al tipo.

    La tableta lleva SIM y también es equipo de cómputo: usa ambas.
    """
    if datos.tipo in (TipoActivo.COMPUTO, TipoActivo.TABLETA) and datos.computo:
        destino = activo.detalle_computo or DetalleComputo(activo_id=activo.id)
        for campo, valor in datos.computo.model_dump().items():
            setattr(destino, campo, valor)
        activo.detalle_computo = destino

    if datos.tipo in (TipoActivo.TELEFONO, TipoActivo.TABLETA) and datos.movil:
        destino = activo.detalle_movil or DetalleMovil(activo_id=activo.id)
        for campo, valor in datos.movil.model_dump().items():
            setattr(destino, campo, valor)
        if datos.tipo == TipoActivo.TELEFONO:
            destino.operador = OPERADOR_TELEFONIA
            destino.esquema = ESQUEMA_TELEFONIA
        activo.detalle_movil = destino

    if (
        datos.tipo in (TipoActivo.LICENCIA_MICROSOFT, TipoActivo.LICENCIA_NORTON)
        and datos.licencia
    ):
        cuenta_nueva = datos.licencia.cuenta_id
        if cuenta_nueva and (
            activo.detalle_licencia is None or activo.detalle_licencia.cuenta_id != cuenta_nueva
        ):
            validar_cupo_familia(db, activo, datos.tipo, cuenta_nueva)
        destino = activo.detalle_licencia or DetalleLicencia(activo_id=activo.id)
        for campo, valor in datos.licencia.model_dump().items():
            setattr(destino, campo, valor)
        activo.detalle_licencia = destino

    elif datos.tipo == TipoActivo.CORREO_WORKSPACE and datos.correo:
        destino = activo.detalle_correo or DetalleCorreo(activo_id=activo.id)
        for campo, valor in datos.correo.model_dump().items():
            setattr(destino, campo, valor)
        activo.detalle_correo = destino


def _crear_activo(db: Session, datos: schemas.ActivoIn) -> Activo:
    if datos.tipo in TIPOS_SIN_VENCIMIENTO:
        datos.fecha_vencimiento = None
    if datos.tipo in TIPOS_LICENCIA:
        datos.fecha_adquisicion = None
    activo = Activo(
        codigo=datos.codigo.strip().upper() if datos.codigo else siguiente_codigo(db, datos.tipo),
        tipo=datos.tipo,
        estatus=datos.estatus,
        descripcion=datos.descripcion,
        fecha_adquisicion=datos.fecha_adquisicion,
        fecha_vencimiento=datos.fecha_vencimiento,
        notas=datos.notas,
    )
    db.add(activo)
    db.flush()
    _aplicar_detalle(db, activo, datos)
    db.flush()
    return activo


def _normalizar_documentos(datos, tipo: TipoActivo) -> None:
    """Deja en NO_APLICA los documentos que el tipo de activo no requiere."""
    if tipo not in TIPOS_CON_RESPONSIVA:
        datos.responsiva_estado = EstadoDocumento.NO_APLICA
        datos.responsiva_fecha_envio = None
        datos.responsiva_fecha_firma = None
        datos.responsiva_destino = None
        datos.responsiva_url = None
        datos.responsiva_archivo = None
        datos.responsiva_nombre = None
    if tipo not in TIPOS_CON_INE:
        datos.ine_estado = EstadoDocumento.NO_APLICA
        datos.ine_url = None
        datos.ine_archivo = None
        datos.ine_nombre = None


def _mensaje_duplicado(exc: IntegrityError) -> str:
    texto = str(exc.orig)
    for campo, etiqueta in (
        ("numero_serie", "el número de serie"),
        ("imei", "el IMEI"),
        ("numero_telefonico", "el número telefónico"),
        ("iccid", "el ICCID"),
        ("direccion", "la dirección de correo"),
        ("codigo", "el código"),
    ):
        if campo in texto:
            return f"{etiqueta} ya está registrado en otro activo"
    return "ya existe un registro con ese valor"


@app.get("/api/activos")
def listar_activos(
    q: str | None = Query(None, max_length=100),
    tipo: TipoActivo | None = None,
    grupo: str | None = None,
    estatus: EstatusActivo | None = None,
    cuenta_id: int | None = None,
    sin_asignar: bool = False,
    sucursal_id: int | None = None,
    departamento_id: int | None = None,
    bodega_id: int | None = None,
    planta_id: int | None = None,
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    consulta = db.query(Activo)
    if grupo:
        if grupo not in GRUPOS:
            raise HTTPException(400, "Grupo desconocido")
        consulta = consulta.filter(Activo.tipo.in_(GRUPOS[grupo]))
    if tipo:
        consulta = consulta.filter(Activo.tipo == tipo)
    if estatus:
        consulta = consulta.filter(Activo.estatus == estatus)
    if cuenta_id:
        consulta = consulta.join(DetalleLicencia).filter(DetalleLicencia.cuenta_id == cuenta_id)
    ubicacion = filtro_ubicacion(sucursal_id, departamento_id, bodega_id, planta_id)
    if ubicacion:
        # Los activos no tienen ubicación propia: se filtran por la persona que los tiene
        consulta = consulta.filter(
            Activo.id.in_(
                db.query(Asignacion.activo_id)
                .join(Empleado, Empleado.id == Asignacion.empleado_id)
                .filter(Asignacion.fecha_devolucion.is_(None), *ubicacion)
            )
        )
    if q:
        patron = patron_busqueda(q)
        consulta = (
            consulta.outerjoin(DetalleComputo)
            .outerjoin(DetalleMovil)
            .outerjoin(DetalleCorreo)
            .outerjoin(DetalleLicencia)
            .filter(
                or_(
                    func.upper(Activo.codigo).like(patron, escape="\\"),
                    func.upper(Activo.descripcion).like(patron, escape="\\"),
                    func.upper(DetalleComputo.numero_serie).like(patron, escape="\\"),
                    func.upper(DetalleMovil.imei).like(patron, escape="\\"),
                    func.upper(DetalleMovil.numero_telefonico).like(patron, escape="\\"),
                    func.upper(DetalleCorreo.direccion).like(patron, escape="\\"),
                    func.upper(DetalleLicencia.correo_licencia).like(patron, escape="\\"),
                    func.upper(DetalleLicencia.identificador_dispositivo).like(patron, escape="\\"),
                )
            )
        )

    # Los detalles y las asignaciones se traen por lotes: una consulta por tabla, no por activo
    activos = (
        consulta.options(
            selectinload(Activo.detalle_computo),
            selectinload(Activo.detalle_movil),
            selectinload(Activo.detalle_licencia).selectinload(DetalleLicencia.cuenta),
            selectinload(Activo.detalle_correo),
        )
        .order_by(Activo.codigo)
        .all()
    )
    vigentes = asignaciones_vigentes(db, [a.id for a in activos])
    salida = [dic_activo(a, db, vigentes) for a in activos]
    if sin_asignar:
        salida = [a for a in salida if a["asignado_a"] is None]
    return salida


@app.get("/api/activos/exportar")
def exportar_activos(
    request: Request,
    grupo: str,
    q: str | None = Query(None, max_length=100),
    tipo: TipoActivo | None = None,
    estatus: EstatusActivo | None = None,
    sin_asignar: bool = False,
    sucursal_id: int | None = None,
    departamento_id: int | None = None,
    bodega_id: int | None = None,
    planta_id: int | None = None,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(usuario_actual),
):
    if grupo not in ("equipos", "tabletas", "telefonia"):
        raise HTTPException(400, "Grupo desconocido")
    activos = listar_activos(
        q, tipo, grupo, estatus, None, sin_asignar,
        sucursal_id, departamento_id, bodega_id, planta_id, db, usuario,
    )
    empleados = {
        e.id: e
        for e in db.query(Empleado).filter(
            Empleado.id.in_([a["asignado_a"]["empleado_id"] for a in activos if a["asignado_a"]])
        )
    }

    def asignacion(a: dict) -> list:
        vig = a["asignado_a"]
        emp = empleados.get(vig["empleado_id"]) if vig else None
        return [
            vig["empleado"] if vig else "",
            vig["codigo_empleado"] if vig else "",
            emp.departamento.nombre if emp and emp.departamento else "",
            emp.sucursal.nombre if emp and emp.sucursal else "",
            vig["fecha_asignacion"] if vig else None,
        ]

    def responsiva(a: dict) -> list:
        vig = a["asignado_a"]
        if not vig:
            return ["", None, "", None, ""]
        return [
            ESTADO_DOC_TEXTO[vig["responsiva_estado"]],
            vig["responsiva_fecha_envio"],
            vig["responsiva_destino"],
            vig["responsiva_fecha_firma"],
            _enlace_documento(request, vig, "responsiva"),
        ]

    columnas_asignacion = ["Asignado a", "Código empleado", "Departamento", "Sucursal", "Asignado desde"]
    columnas_responsiva = ["Responsiva", "Enviada a firmar", "Enviada a", "Firmada el", "Documento responsiva"]

    if grupo == "telefonia":
        filas = []
        for a in activos:
            m = a.get("movil") or {}
            vig = a["asignado_a"]
            filas.append(
                [
                    a["codigo"], m.get("marca"), m.get("modelo"), m.get("numero_telefonico"),
                    m.get("imei"), m.get("iccid"), m.get("operador"), m.get("precio_plan"),
                    m.get("esquema"), "Sí" if m.get("esim") else "No", m.get("cuenta_operador"),
                    a["estatus"], *asignacion(a), *responsiva(a),
                    ESTADO_DOC_TEXTO[vig["ine_estado"]] if vig else "",
                    _enlace_documento(request, vig, "ine"), a["notas"],
                ]
            )
        return _descarga_excel(
            "Telefonía",
            ["Código", "Marca", "Modelo", "Teléfono", "IMEI", "ICCID", "Operador", "Plan",
             "Esquema", "eSIM", "Cuenta operador", "Estatus", *columnas_asignacion,
             *columnas_responsiva, "INE", "Documento INE", "Notas"],
            filas,
            [12, 14, 22, 14, 18, 22, 12, 10, 16, 7, 16, 13, 30, 14, 20, 16, 13,
             20, 14, 20, 13, 30, 20, 30, 30],
            "telefonia",
        )

    filas = []
    for a in activos:
        c = a.get("computo") or {}
        filas.append(
            [
                a["codigo"], c.get("marca"), c.get("modelo"), c.get("numero_serie"),
                c.get("procesador"), c.get("ram_gb"), c.get("almacenamiento"),
                c.get("sistema_operativo"), a["fecha_adquisicion"], a["estatus"],
                *asignacion(a), *responsiva(a), a["notas"],
            ]
        )
    return _descarga_excel(
        "Tablets" if grupo == "tabletas" else "Equipos de cómputo",
        ["Código", "Marca", "Modelo", "Número de serie", "Procesador", "RAM (GB)",
         "Almacenamiento", "Sistema operativo", "Fecha de adquisición", "Estatus",
         *columnas_asignacion, *columnas_responsiva, "Notas"],
        filas,
        [12, 14, 22, 20, 26, 9, 15, 18, 14, 13, 30, 14, 20, 16, 13, 20, 14, 20, 13, 30, 30],
        grupo,
    )


@app.get("/api/activos/{activo_id}")
def ver_activo(activo_id: int, db: Session = Depends(obtener_sesion), _=Depends(usuario_actual)):
    activo = db.get(Activo, activo_id)
    if activo is None:
        raise HTTPException(404, "Activo no encontrado")
    datos = dic_activo(activo, db)
    datos["historial"] = [
        dic_asignacion(a)
        for a in db.query(Asignacion)
        .filter(Asignacion.activo_id == activo_id)
        .order_by(Asignacion.fecha_asignacion.desc())
        .all()
    ]
    return datos


@app.post("/api/activos")
@un_alta_a_la_vez
def crear_activo(
    datos: schemas.ActivoIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    try:
        activo = _crear_activo(db, datos)
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, f"Dato duplicado: {_mensaje_duplicado(exc)}")
    registrar(db, usuario, "CREAR", "activo", activo.id, activo.codigo)
    db.commit()
    return dic_activo(activo, db)


@app.put("/api/activos/{activo_id}")
def editar_activo(
    activo_id: int,
    datos: schemas.ActivoIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    activo = db.get(Activo, activo_id)
    if activo is None:
        raise HTTPException(404, "Activo no encontrado")

    activo.estatus = datos.estatus
    activo.descripcion = datos.descripcion
    activo.fecha_adquisicion = None if activo.tipo in TIPOS_LICENCIA else datos.fecha_adquisicion
    activo.fecha_vencimiento = (
        None if activo.tipo in TIPOS_SIN_VENCIMIENTO else datos.fecha_vencimiento
    )
    activo.notas = datos.notas
    if datos.codigo:
        activo.codigo = datos.codigo.strip().upper()

    try:
        _aplicar_detalle(db, activo, datos)
        db.flush()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, f"Dato duplicado: {_mensaje_duplicado(exc)}")
    registrar(db, usuario, "EDITAR", "activo", activo.id, activo.codigo)
    db.commit()
    return dic_activo(activo, db)


@app.delete("/api/activos/{activo_id}")
def eliminar_activo(
    activo_id: int,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    activo = db.get(Activo, activo_id)
    if activo is None:
        raise HTTPException(404, "Activo no encontrado")

    codigo = activo.codigo
    # Se borra también su historial de asignaciones: la interfaz avisa antes.
    db.query(Asignacion).filter(Asignacion.activo_id == activo_id).delete()
    db.delete(activo)
    registrar(db, usuario, "ELIMINAR", "activo", activo_id, codigo)
    db.commit()
    return {"ok": True}


COLUMNAS_REPORTE_CORREO = {
    "direccion": "email address",
    "ultimo_acceso": "last sign in",
    "almacenamiento": "email usage",
}
# Columnas C, E y F del reporte de Google Workspace
POSICION_REPORTE_CORREO = {"direccion": 2, "ultimo_acceso": 4, "almacenamiento": 5}
UNIDADES_GB = {"KB": 1 / 1024**2, "MB": 1 / 1024, "GB": 1, "TB": 1024}


def _leer_ultimo_acceso(valor) -> tuple[bool, dt.datetime | None]:
    """(se pudo leer, fecha). Las cuentas que nunca han entrado dan (True, None)."""
    if isinstance(valor, dt.datetime):
        return True, valor
    texto = str(valor or "").strip()
    if not texto or texto.lower().startswith("never"):
        return True, None
    for formato in ("%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%Y/%m/%d", "%Y-%m-%d"):
        try:
            return True, dt.datetime.strptime(texto, formato)
        except ValueError:
            pass
    return False, None


def _leer_almacenamiento(valor) -> float | None:
    if isinstance(valor, (int, float)):
        return float(valor)
    coincidencia = re.match(r"^\s*([\d.,]+)\s*([KMGT]B)?\s*$", str(valor or ""), re.IGNORECASE)
    if not coincidencia:
        return None
    numero = float(coincidencia.group(1).replace(",", ""))
    return round(numero * UNIDADES_GB[(coincidencia.group(2) or "GB").upper()], 2)


@app.post("/api/correos/importar-actividad")
async def importar_actividad_correos(
    archivo: UploadFile = File(...),
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    """Carga el reporte de usuarios de Google Workspace: último acceso y almacenamiento."""
    from openpyxl import load_workbook

    if Path(archivo.filename or "").suffix.lower() != ".xlsx":
        raise HTTPException(400, "El reporte debe ser un archivo .xlsx")
    contenido = await archivo.read()
    if len(contenido) > TAMANO_MAXIMO:
        raise HTTPException(400, "El archivo supera los 10 MB")
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo de Excel")

    filas = libro.active.iter_rows(values_only=True)
    encabezado = [str(c or "").strip().lower() for c in next(filas, [])]
    indices = {}
    for clave, prefijo in COLUMNAS_REPORTE_CORREO.items():
        indice = next((i for i, c in enumerate(encabezado) if c.startswith(prefijo)), None)
        indices[clave] = POSICION_REPORTE_CORREO[clave] if indice is None else indice

    cuentas = {d.direccion.strip().lower(): d for d in db.query(DetalleCorreo).all()}
    momento = dt.datetime.now(dt.timezone.utc)
    actualizados = 0
    no_registrados: list[str] = []
    inactivos_no_registrados: list[dict] = []
    ahora = ahora_local()
    for fila in filas:
        if not fila or len(fila) <= max(indices.values()):
            continue
        direccion = str(fila[indices["direccion"]] or "").strip().lower()
        if not direccion:
            continue
        leido, ultimo = _leer_ultimo_acceso(fila[indices["ultimo_acceso"]])
        cuenta = cuentas.get(direccion)
        if cuenta is None:
            no_registrados.append(direccion)
            dias = (ahora - ultimo).days if ultimo else None
            if leido and (dias is None or dias >= DIAS_DESUSO_CORREO):
                inactivos_no_registrados.append(
                    {
                        "direccion": direccion,
                        "ultimo_acceso": ultimo,
                        "dias_sin_acceso": dias,
                        "almacenamiento_gb": _leer_almacenamiento(fila[indices["almacenamiento"]]),
                    }
                )
            continue
        if leido:
            cuenta.ultimo_acceso = ultimo
        almacenamiento = _leer_almacenamiento(fila[indices["almacenamiento"]])
        if almacenamiento is not None:
            cuenta.almacenamiento_gb = almacenamiento
        cuenta.actividad_importada_en = momento
        actualizados += 1
    libro.close()

    registrar(
        db, usuario, "EDITAR", "correos", None,
        f"actividad importada de {archivo.filename}: {actualizados} cuenta(s)",
    )
    db.commit()
    return {
        "actualizados": actualizados,
        "no_registrados": sorted(no_registrados),
        "inactivos_no_registrados": sorted(
            inactivos_no_registrados,
            key=lambda c: -1 if c["dias_sin_acceso"] is None else -c["dias_sin_acceso"],
        ),
        "en_desuso": sum(1 for d in cuentas.values() if correo_en_desuso(d)),
    }


@app.post("/api/altas/equipo")
@un_alta_a_la_vez
def alta_equipo(
    datos: schemas.AltaEquipoIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    """Da de alta una PC o una tableta y, opcionalmente, su correo y sus licencias.

    Todo se asigna al mismo empleado. Cada pieza queda como activo propio,
    editable y dable de baja por separado.
    """
    if datos.equipo.tipo not in (TipoActivo.COMPUTO, TipoActivo.TABLETA):
        raise HTTPException(400, "Esta alta es sólo para equipos de cómputo y tabletas")

    empleado = None
    if datos.empleado_id:
        empleado = db.get(Empleado, datos.empleado_id)
        if empleado is None:
            raise HTTPException(404, "Empleado no encontrado")
        if empleado.estatus == EstatusEmpleado.BAJA:
            raise HTTPException(409, "No se puede asignar a un empleado dado de baja")

    creados: list[Activo] = []
    try:
        creados.append(_crear_activo(db, datos.equipo))

        comunes = {"estatus": EstatusActivo.DISPONIBLE}

        if datos.correo and datos.correo.direccion:
            creados.append(
                _crear_activo(
                    db,
                    schemas.ActivoIn(
                        tipo=TipoActivo.CORREO_WORKSPACE,
                        descripcion=datos.correo.direccion,
                        correo=datos.correo,
                        **comunes,
                    ),
                )
            )

        for licencia, tipo, etiqueta in (
            (datos.licencia_microsoft, TipoActivo.LICENCIA_MICROSOFT, "Microsoft 365"),
            (datos.licencia_norton, TipoActivo.LICENCIA_NORTON, "Norton 360"),
        ):
            if licencia and (
                licencia.cuenta_id
                or licencia.correo_licencia
                or licencia.identificador_dispositivo
            ):
                creados.append(
                    _crear_activo(
                        db,
                        schemas.ActivoIn(
                            tipo=tipo,
                            descripcion=licencia.correo_licencia or etiqueta,
                            licencia=licencia,
                            **comunes,
                        ),
                    )
                )
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, f"Dato duplicado: {_mensaje_duplicado(exc)}")

    asignados = 0
    if empleado:
        for activo in creados:
            lleva = activo.tipo in TIPOS_CON_RESPONSIVA
            lleva_ine = activo.tipo in TIPOS_CON_INE
            db.add(
                Asignacion(
                    activo_id=activo.id,
                    empleado_id=empleado.id,
                    fecha_asignacion=datos.fecha_asignacion or hoy_local(),
                    responsiva_estado=datos.responsiva_estado
                    if lleva
                    else EstadoDocumento.NO_APLICA,
                    responsiva_fecha_envio=datos.responsiva_fecha_envio if lleva else None,
                    responsiva_fecha_firma=datos.responsiva_fecha_firma if lleva else None,
                    responsiva_destino=datos.responsiva_destino if lleva else None,
                    responsiva_url=datos.responsiva_url if lleva else None,
                    responsiva_archivo=datos.responsiva_archivo if lleva else None,
                    responsiva_nombre=datos.responsiva_nombre if lleva else None,
                    ine_estado=datos.ine_estado if lleva_ine else EstadoDocumento.NO_APLICA,
                    ine_url=datos.ine_url if lleva_ine else None,
                    ine_archivo=datos.ine_archivo if lleva_ine else None,
                    ine_nombre=datos.ine_nombre if lleva_ine else None,
                )
            )
            activo.estatus = EstatusActivo.ASIGNADO
            asignados += 1
        db.flush()

    registrar(
        db,
        usuario,
        "CREAR",
        "alta_equipo",
        creados[0].id,
        f"{', '.join(a.codigo for a in creados)}"
        + (f" → {empleado.nombre_completo}" if empleado else ""),
    )
    db.commit()
    return {
        "creados": [dic_activo(a, db) for a in creados],
        "asignados": asignados,
    }


# --------------------------------------------------------------------------
# Asignaciones
# --------------------------------------------------------------------------


@app.post("/api/asignaciones")
def asignar(
    datos: schemas.AsignacionIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    activo = db.get(Activo, datos.activo_id)
    empleado = db.get(Empleado, datos.empleado_id)
    if activo is None:
        raise HTTPException(404, "Activo no encontrado")
    if empleado is None:
        raise HTTPException(404, "Empleado no encontrado")
    if asignacion_vigente(db, activo.id):
        raise HTTPException(
            409, f"{activo.codigo} ya está asignado. Regístralo como devuelto primero."
        )
    if empleado.estatus == EstatusEmpleado.BAJA:
        raise HTTPException(409, "No se puede asignar a un empleado dado de baja")

    _normalizar_documentos(datos, activo.tipo)
    asignacion = Asignacion(
        activo_id=datos.activo_id,
        empleado_id=datos.empleado_id,
        fecha_asignacion=datos.fecha_asignacion or hoy_local(),
        responsiva_estado=datos.responsiva_estado,
        responsiva_fecha_envio=datos.responsiva_fecha_envio,
        responsiva_fecha_firma=datos.responsiva_fecha_firma,
        responsiva_destino=datos.responsiva_destino,
        responsiva_url=datos.responsiva_url,
        responsiva_archivo=datos.responsiva_archivo,
        responsiva_nombre=datos.responsiva_nombre,
        ine_estado=datos.ine_estado,
        ine_url=datos.ine_url,
        ine_archivo=datos.ine_archivo,
        ine_nombre=datos.ine_nombre,
        notas=datos.notas,
    )
    db.add(asignacion)
    activo.estatus = EstatusActivo.ASIGNADO
    db.flush()
    registrar(
        db,
        usuario,
        "ASIGNAR",
        "asignacion",
        asignacion.id,
        f"{activo.codigo} → {empleado.nombre_completo}",
    )
    db.commit()
    return dic_asignacion(asignacion)


@app.put("/api/asignaciones/{asignacion_id}")
def editar_asignacion(
    asignacion_id: int,
    datos: schemas.AsignacionEditarIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    asignacion = db.get(Asignacion, asignacion_id)
    if asignacion is None:
        raise HTTPException(404, "Asignación no encontrada")
    _normalizar_documentos(datos, asignacion.activo.tipo)
    for campo, valor in datos.model_dump().items():
        setattr(asignacion, campo, valor)
    registrar(
        db, usuario, "EDITAR", "asignacion", asignacion.id, f"documentación de {asignacion.activo.codigo}"
    )
    db.commit()
    return dic_asignacion(asignacion)


@app.post("/api/asignaciones/{asignacion_id}/devolver")
def devolver(
    asignacion_id: int,
    datos: schemas.DevolucionIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    asignacion = db.get(Asignacion, asignacion_id)
    if asignacion is None:
        raise HTTPException(404, "Asignación no encontrada")
    if asignacion.fecha_devolucion is not None:
        raise HTTPException(409, "Esta asignación ya fue cerrada")

    asignacion.fecha_devolucion = datos.fecha_devolucion or hoy_local()
    asignacion.motivo_devolucion = datos.motivo_devolucion
    if datos.notas:
        asignacion.notas = f"{asignacion.notas or ''}\n{datos.notas}".strip()
    asignacion.activo.estatus = datos.nuevo_estatus_activo

    registrar(
        db,
        usuario,
        "DEVOLVER",
        "asignacion",
        asignacion.id,
        f"{asignacion.activo.codigo} ← {asignacion.empleado.nombre_completo}",
    )
    db.commit()
    return dic_asignacion(asignacion)


# --------------------------------------------------------------------------
# Tablero y alertas
# --------------------------------------------------------------------------


@app.get("/api/tablero")
def tablero(db: Session = Depends(obtener_sesion), _=Depends(usuario_actual)):
    por_tipo = dict(db.query(Activo.tipo, func.count(Activo.id)).group_by(Activo.tipo).all())
    fisico = Activo.tipo.in_(TIPOS_FISICOS)
    por_estatus = dict(
        db.query(Activo.estatus, func.count(Activo.id))
        .filter(fisico)
        .group_by(Activo.estatus)
        .all()
    )
    asignados = (
        db.query(func.count(Asignacion.id))
        .join(Activo, Activo.id == Asignacion.activo_id)
        .filter(Asignacion.fecha_devolucion.is_(None), fisico)
        .scalar()
    )
    def contar_documento(estado):
        return (
            db.query(func.count(Asignacion.id))
            .filter(
                Asignacion.fecha_devolucion.is_(None),
                Asignacion.responsiva_estado == estado,
            )
            .scalar()
        )

    return {
        "empleados_activos": db.query(func.count(Empleado.id))
        .filter(Empleado.estatus == EstatusEmpleado.ACTIVO)
        .scalar(),
        "empleados_baja": db.query(func.count(Empleado.id))
        .filter(Empleado.estatus == EstatusEmpleado.BAJA)
        .scalar(),
        "total_activos": db.query(func.count(Activo.id)).filter(fisico).scalar(),
        "activos_asignados": asignados,
        "activos_disponibles": db.query(func.count(Activo.id))
        .filter(fisico, Activo.estatus == EstatusActivo.DISPONIBLE)
        .scalar(),
        "total_servicios": db.query(func.count(Activo.id)).filter(~fisico).scalar(),
        "responsivas_pendientes": contar_documento(EstadoDocumento.PENDIENTE),
        "responsivas_en_espera": contar_documento(EstadoDocumento.ENVIADA),
        "responsivas_firmadas": contar_documento(EstadoDocumento.FIRMADA),
        "notas_abiertas": db.query(func.count(Nota.id))
        .filter(Nota.estado != EstadoNota.RESUELTA)
        .scalar(),
        "por_tipo": {t.value: n for t, n in por_tipo.items()},
        "por_estatus": {e.value: n for e, n in por_estatus.items()},
        "por_departamento": activos_por_departamento(db),
        "vencimientos_licencias": vencimientos_licencias(db),
    }


def activos_por_departamento(db: Session) -> list[dict]:
    """Activos físicos vigentes según el departamento de quien los tiene."""
    filas = (
        db.query(Departamento.nombre, func.count(Asignacion.id))
        .select_from(Asignacion)
        .join(Activo, Activo.id == Asignacion.activo_id)
        .join(Empleado, Empleado.id == Asignacion.empleado_id)
        .outerjoin(Departamento, Departamento.id == Empleado.departamento_id)
        .filter(Asignacion.fecha_devolucion.is_(None), Activo.tipo.in_(TIPOS_FISICOS))
        .group_by(Departamento.nombre)
        .all()
    )
    salida = [{"departamento": nombre or "SIN DEPARTAMENTO", "activos": n} for nombre, n in filas]
    return sorted(salida, key=lambda f: -f["activos"])


def vencimientos_licencias(db: Session) -> list[dict]:
    """Renovación de cada familia de Microsoft 365 y Norton, y de licencias sueltas con fecha."""
    hoy = hoy_local()
    salida = []
    for c in db.query(CuentaLicencia).filter(CuentaLicencia.activo.is_(True)):
        tipos = {d.activo.tipo for d in c.licencias}
        producto = (
            "NORTON"
            if TipoActivo.LICENCIA_NORTON in tipos or "norton" in (c.producto or "").lower()
            else "MICROSOFT"
        )
        salida.append(
            {
                "producto": producto,
                "nombre": c.nombre,
                "correo": c.correo_cuenta,
                "fecha": c.fecha_renovacion,
                "dias": (c.fecha_renovacion - hoy).days if c.fecha_renovacion else None,
                "licencias": len(c.licencias),
                "cuenta_id": c.id,
            }
        )
    for a in (
        db.query(Activo)
        .join(DetalleLicencia)
        .filter(DetalleLicencia.cuenta_id.is_(None), Activo.fecha_vencimiento.isnot(None))
    ):
        salida.append(
            {
                "producto": "NORTON" if a.tipo == TipoActivo.LICENCIA_NORTON else "MICROSOFT",
                "nombre": a.codigo,
                "correo": a.detalle_licencia.correo_licencia or a.detalle_licencia.identificador_dispositivo,
                "fecha": a.fecha_vencimiento,
                "dias": (a.fecha_vencimiento - hoy).days,
                "licencias": 1,
                "activo_id": a.id,
            }
        )
    return sorted(salida, key=lambda v: (v["fecha"] is None, v["fecha"] or hoy))


@app.get("/api/notificaciones")
def notificaciones(db: Session = Depends(obtener_sesion), _=Depends(usuario_actual)):
    """Responsivas que todavía no regresan firmadas, de la más antigua a la más reciente."""
    pendientes = (
        db.query(Asignacion)
        .filter(
            Asignacion.fecha_devolucion.is_(None),
            Asignacion.responsiva_estado.in_(
                [EstadoDocumento.PENDIENTE, EstadoDocumento.ENVIADA]
            ),
        )
        .order_by(Asignacion.fecha_asignacion)
        .all()
    )
    salida = []
    for a in pendientes:
        datos = dic_asignacion(a)
        datos["atrasada"] = bool(
            datos["dias_en_espera"] is not None
            and datos["dias_en_espera"] > DIAS_ESPERA_RESPONSIVA
        )
        salida.append(datos)
    return salida


@app.get("/api/alertas")
def alertas(
    dias: int = Query(60, ge=1, le=365),
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    hoy = hoy_local()
    resultado: list[dict] = []

    # Licencias: una alerta por familia (o por licencia suelta) a 60 días o menos
    for v in vencimientos_licencias(db):
        if v["dias"] is None or v["dias"] > dias:
            continue
        producto = "Norton 360" if v["producto"] == "NORTON" else "Microsoft 365"
        if v["dias"] < 0:
            cuando = f"venció hace {-v['dias']} día(s)"
        elif v["dias"] == 0:
            cuando = "vence hoy"
        else:
            cuando = f"vence en {v['dias']} día(s)"
        resultado.append(
            {
                "severidad": "critica" if v["dias"] <= 30 else "media",
                "tipo": "licencia_por_vencer",
                "titulo": f"{v['nombre']} ({producto}) {cuando}",
                "detalle": f"Renovación el {v['fecha']:%d/%m/%Y} · {v['licencias']} licencia(s)"
                + (f" · {v['correo']}" if v["correo"] else ""),
                "dias": v["dias"],
                "producto": v["producto"].lower(),
                "cuenta_id": v.get("cuenta_id"),
                "activo_id": v.get("activo_id"),
            }
        )

    for c in db.query(CuentaLicencia).filter_by(activo=True).all():
        usadas = db.query(DetalleLicencia).filter_by(cuenta_id=c.id).count()
        if usadas > c.ranuras_totales:
            resultado.append(
                {
                    "severidad": "critica",
                    "tipo": "sobreasignacion",
                    "titulo": f"{c.nombre} tiene {usadas} usos sobre {c.ranuras_totales} ranuras",
                    "detalle": "Se están usando más licencias de las contratadas",
                    "cuenta_id": c.id,
                    "producto": "norton" if "norton" in (c.producto or "").lower() else "microsoft",
                }
            )

    for d in (
        db.query(DetalleCorreo)
        .join(Activo, Activo.id == DetalleCorreo.activo_id)
        .filter(Activo.estatus != EstatusActivo.BAJA)
        .order_by(DetalleCorreo.ultimo_acceso)
    ):
        if not correo_en_desuso(d):
            continue
        vig = asignacion_vigente(db, d.activo_id)
        partes = [
            f"Último acceso {d.ultimo_acceso:%d/%m/%Y}" if d.ultimo_acceso else "Sin accesos registrados",
            f"{float(d.almacenamiento_gb):.2f} GB" if d.almacenamiento_gb is not None else None,
            vig.empleado.nombre_completo if vig else "Sin asignar",
        ]
        resultado.append(
            {
                "severidad": "media",
                "tipo": "correo_sin_uso",
                "titulo": f"{d.direccion} nunca ha iniciado sesión"
                if d.ultimo_acceso is None
                else f"{d.direccion} lleva {dias_sin_acceso(d)} días sin iniciar sesión",
                "detalle": " · ".join(p for p in partes if p),
                "activo_id": d.activo_id,
            }
        )

    baja_con_activos = (
        db.query(Empleado)
        .join(Asignacion, Asignacion.empleado_id == Empleado.id)
        .filter(Empleado.estatus == EstatusEmpleado.BAJA, Asignacion.fecha_devolucion.is_(None))
        .distinct()
        .all()
    )
    for e in baja_con_activos:
        pendientes = (
            db.query(func.count(Asignacion.id))
            .filter(Asignacion.empleado_id == e.id, Asignacion.fecha_devolucion.is_(None))
            .scalar()
        )
        resultado.append(
            {
                "severidad": "critica",
                "tipo": "baja_con_activos",
                "titulo": f"{e.nombre_completo} está dado de baja y conserva {pendientes} activo(s)",
                "detalle": "Registrar la devolución de los equipos",
                "empleado_id": e.id,
            }
        )

    pendientes = (
        db.query(func.count(Asignacion.id))
        .filter(
            Asignacion.fecha_devolucion.is_(None),
            Asignacion.responsiva_estado == EstadoDocumento.PENDIENTE,
        )
        .scalar()
    )
    if pendientes:
        resultado.append(
            {
                "severidad": "media",
                "tipo": "responsiva",
                "titulo": f"{pendientes} activo(s) entregados sin responsiva",
                "detalle": "Todavía no se envían a firmar",
            }
        )

    # Responsivas que llevan demasiado tiempo fuera esperando firma
    atrasadas = (
        db.query(Asignacion)
        .filter(
            Asignacion.fecha_devolucion.is_(None),
            Asignacion.responsiva_estado == EstadoDocumento.ENVIADA,
            Asignacion.responsiva_fecha_envio.isnot(None),
            Asignacion.responsiva_fecha_envio <= hoy - dt.timedelta(days=DIAS_ESPERA_RESPONSIVA),
        )
        .all()
    )
    for a in atrasadas:
        dias = (hoy - a.responsiva_fecha_envio).days
        resultado.append(
            {
                "severidad": "media",
                "tipo": "responsiva_atrasada",
                "titulo": f"Responsiva de {a.activo.codigo} lleva {dias} días esperando firma",
                "detalle": a.empleado.nombre_completo
                + (f" · enviada a {a.responsiva_destino}" if a.responsiva_destino else ""),
                "activo_id": a.activo_id,
                "asignacion_id": a.id,
            }
        )

    # Auditorias semanales sin cerrar
    for aud in (
        db.query(Auditoria)
        .filter(Auditoria.estado == EstadoAuditoria.EN_PROCESO, Auditoria.fecha_fin < hoy)
        .all()
    ):
        abiertos = (
            db.query(func.count(ItemAuditoria.id))
            .filter(
                ItemAuditoria.auditoria_id == aud.id,
                ItemAuditoria.estado == EstadoItem.PENDIENTE,
            )
            .scalar()
        )
        resultado.append(
            {
                "severidad": "media",
                "tipo": "auditoria",
                "titulo": f"Auditoría de {aud.tipo.value.lower()} de la semana del "
                f"{aud.fecha_inicio:%d/%m} sigue abierta",
                "detalle": f"{abiertos} punto(s) sin revisar"
                + (f" · {aud.analista_nombre}" if aud.analista_nombre else ""),
                "auditoria_id": aud.id,
            }
        )

    corte = hoy - dt.timedelta(days=4 * 365)
    viejos = (
        db.query(func.count(Activo.id))
        .filter(
            Activo.tipo == TipoActivo.COMPUTO,
            Activo.fecha_adquisicion.isnot(None),
            Activo.fecha_adquisicion <= corte,
            Activo.estatus != EstatusActivo.BAJA,
        )
        .scalar()
    )
    if viejos:
        resultado.append(
            {
                "severidad": "baja",
                "tipo": "reemplazo",
                "titulo": f"{viejos} equipo(s) de cómputo superan 4 años de antigüedad",
                "detalle": "Candidatos a reemplazo",
            }
        )

    orden = {"critica": 0, "media": 1, "baja": 2}
    resultado.sort(key=lambda a: orden.get(a["severidad"], 9))
    return resultado


# --------------------------------------------------------------------------
# Reportes
# --------------------------------------------------------------------------


@app.get("/api/reportes/inventario")
def reporte_inventario(
    agrupar_por: str = Query("sucursal", pattern="^(sucursal|departamento|tipo|estatus)$"),
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    fisico = Activo.tipo.in_(TIPOS_FISICOS)
    if agrupar_por == "tipo":
        datos = (
            db.query(Activo.tipo, func.count(Activo.id)).filter(fisico).group_by(Activo.tipo).all()
        )
        filas = [{"grupo": t.value, "activos": n} for t, n in datos]

    elif agrupar_por == "estatus":
        datos = (
            db.query(Activo.estatus, func.count(Activo.id))
            .filter(fisico)
            .group_by(Activo.estatus)
            .all()
        )
        filas = [{"grupo": e.value, "activos": n} for e, n in datos]

    else:
        # La ubicación es de la persona: el activo se cuenta donde está quien lo tiene
        Catalogo, campo = (
            (Sucursal, Empleado.sucursal_id)
            if agrupar_por == "sucursal"
            else (Departamento, Empleado.departamento_id)
        )
        datos = (
            db.query(Catalogo.nombre, func.count(Asignacion.id))
            .select_from(Catalogo)
            .join(Empleado, campo == Catalogo.id)
            .join(
                Asignacion,
                (Asignacion.empleado_id == Empleado.id)
                & (Asignacion.fecha_devolucion.is_(None)),
            )
            .join(Activo, (Activo.id == Asignacion.activo_id) & fisico)
            .group_by(Catalogo.nombre)
            .all()
        )
        filas = [{"grupo": nombre, "activos": n} for nombre, n in datos]

    filas = [f for f in filas if f["activos"]]
    filas.sort(key=lambda f: -f["activos"])
    return filas


@app.get("/api/reportes/exportar")
def exportar_excel(db: Session = Depends(obtener_sesion), _=Depends(usuario_actual)):
    from openpyxl import Workbook

    wb = Workbook()

    hoja = wb.active
    hoja.title = "Activos"
    hoja.append(
        [
            "Código", "Tipo", "Estatus", "Descripción", "Serie/IMEI/Correo",
            "Familia de licencia", "Asignado a", "Código empleado", "Departamento",
            "Sucursal", "Locación", "Fecha adquisición", "Vencimiento", "Responsiva",
        ]
    )
    for a in db.query(Activo).order_by(Activo.tipo, Activo.codigo).all():
        vig = asignacion_vigente(db, a.id)
        identificador = ""
        if a.detalle_computo:
            identificador = a.detalle_computo.numero_serie or ""
        elif a.detalle_movil:
            identificador = a.detalle_movil.imei or a.detalle_movil.numero_telefonico or ""
        elif a.detalle_correo:
            identificador = a.detalle_correo.direccion
        elif a.detalle_licencia:
            identificador = a.detalle_licencia.correo_licencia or ""

        empleado = vig.empleado if vig else None
        hoja.append(
            [
                a.codigo, a.tipo.value, a.estatus.value, a.descripcion, identificador,
                a.detalle_licencia.cuenta.nombre
                if a.detalle_licencia and a.detalle_licencia.cuenta
                else "",
                empleado.nombre_completo if empleado else "",
                empleado.codigo if empleado else "",
                empleado.departamento.nombre if empleado and empleado.departamento else "",
                empleado.sucursal.nombre if empleado and empleado.sucursal else "",
                lugar_trabajo(empleado) if empleado else "",
                a.fecha_adquisicion, a.fecha_vencimiento,
                vig.responsiva_estado.value if vig else "",
            ]
        )

    hoja2 = wb.create_sheet("Empleados")
    hoja2.append(
        ["Código", "Nombre completo", "Correos asignados", "Puesto", "Sucursal",
         "Departamento", "Locación", "Estatus", "Activos vigentes"]
    )
    todos = db.query(Empleado).order_by(Empleado.nombre_completo).all()
    correos = correos_asignados(db, [e.id for e in todos])
    for e in todos:
        total = (
            db.query(func.count(Asignacion.id))
            .join(Activo, Activo.id == Asignacion.activo_id)
            .filter(
                Asignacion.empleado_id == e.id,
                Asignacion.fecha_devolucion.is_(None),
                Activo.tipo.in_(TIPOS_FISICOS),
            )
            .scalar()
        )
        hoja2.append(
            [
                e.codigo, e.nombre_completo, ", ".join(correos.get(e.id, [])), e.puesto,
                e.sucursal.nombre if e.sucursal else "",
                e.departamento.nombre if e.departamento else "",
                lugar_trabajo(e),
                e.estatus.value, total,
            ]
        )

    hoja3 = wb.create_sheet("Historial")
    hoja3.append(["Activo", "Tipo", "Empleado", "Desde", "Hasta", "Motivo", "Responsiva"])
    for a in db.query(Asignacion).order_by(Asignacion.fecha_asignacion.desc()).all():
        hoja3.append(
            [
                a.activo.codigo, a.activo.tipo.value, a.empleado.nombre_completo,
                a.fecha_asignacion, a.fecha_devolucion, a.motivo_devolucion,
                a.responsiva_estado.value,
            ]
        )

    hoja4 = wb.create_sheet("Impresoras")
    hoja4.append(
        ["Número de serie", "Propietario", "Locación", "Departamentos",
         "Número de serie del proveedor", "Identificador del proveedor"]
    )
    for i in db.query(Impresora).order_by(Impresora.numero_serie).all():
        hoja4.append(
            [
                i.numero_serie, i.propietario.value, i.locacion,
                ", ".join(d.nombre for d in i.departamentos),
                i.numero_serie_proveedor, i.identificador_proveedor,
            ]
        )

    _neutralizar_formulas(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nombre = f"inventario_ti_{hoy_local():%Y%m%d}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


# --------------------------------------------------------------------------
# Usuarios del sistema y bitacora
# --------------------------------------------------------------------------


def id_admin_principal(db: Session) -> int | None:
    """El primer administrador creado. Nadie más puede modificarlo ni otorgar el rol ADMIN."""
    return (
        db.query(func.min(Usuario.id)).filter(Usuario.rol == RolUsuario.ADMIN).scalar()
    )


@app.get("/api/usuarios")
def listar_usuarios(db: Session = Depends(obtener_sesion), _=Depends(requiere_admin)):
    principal = id_admin_principal(db)
    return [
        {
            "id": u.id,
            "username": u.username,
            "nombre": u.nombre,
            "rol": u.rol.value,
            "activo": u.activo,
            "ultimo_acceso": u.ultimo_acceso,
            "creado_en": u.creado_en,
            "principal": u.id == principal,
            "foto": u.foto,
            "correo": u.correo,
            "puesto": u.puesto,
        }
        for u in db.query(Usuario).order_by(Usuario.id)
    ]


@app.post("/api/usuarios")
def crear_usuario(
    datos: schemas.UsuarioIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_admin),
):
    if datos.rol == RolUsuario.ADMIN and usuario.id != id_admin_principal(db):
        raise HTTPException(403, "Sólo el administrador principal puede crear administradores")
    nuevo = Usuario(
        username=datos.username.strip().lower(),
        nombre=datos.nombre,
        password_hash=hashear_password(datos.password),
        rol=datos.rol,
    )
    db.add(nuevo)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Ese nombre de usuario ya existe")
    registrar(db, usuario, "CREAR", "usuario", nuevo.id, nuevo.username)
    db.commit()
    return {"id": nuevo.id, "username": nuevo.username}


@app.put("/api/usuarios/{usuario_id}")
def editar_usuario(
    usuario_id: int,
    datos: schemas.UsuarioEditarIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_admin),
):
    objetivo = db.get(Usuario, usuario_id)
    if objetivo is None:
        raise HTTPException(404, "Usuario no encontrado")

    principal = id_admin_principal(db)
    soy_principal = usuario.id == principal
    if objetivo.id == usuario.id:
        raise HTTPException(400, "Tu propia cuenta se administra desde «Cambiar contraseña»")
    if objetivo.id == principal:
        raise HTTPException(403, "El administrador principal no puede modificarse")
    if not soy_principal and (objetivo.rol == RolUsuario.ADMIN or datos.rol == RolUsuario.ADMIN):
        raise HTTPException(403, "Sólo el administrador principal gestiona administradores")

    objetivo.nombre = datos.nombre
    objetivo.rol = datos.rol
    objetivo.activo = datos.activo
    if datos.password:
        objetivo.password_hash = hashear_password(datos.password)
    if not datos.activo or datos.password:
        db.query(Sesion).filter(Sesion.usuario_id == objetivo.id).delete()

    registrar(
        db, usuario, "EDITAR", "usuario", objetivo.id,
        f"{objetivo.username} · {datos.rol.value} · {'activo' if datos.activo else 'inactivo'}",
    )
    db.commit()
    return {"ok": True}


@app.get("/api/bitacora")
def ver_bitacora(
    limite: int = Query(100, ge=1, le=500),
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    return [
        {
            "id": b.id,
            "fecha": b.fecha,
            "usuario": b.usuario_nombre,
            "accion": b.accion,
            "entidad": b.entidad,
            "entidad_id": b.entidad_id,
            "detalle": b.detalle,
        }
        for b in db.query(Bitacora).order_by(Bitacora.fecha.desc()).limit(limite)
    ]



# --------------------------------------------------------------------------
# CRUD de catalogos
# --------------------------------------------------------------------------

MODELOS_CATALOGO = {
    "sucursales": Sucursal,
    "departamentos": Departamento,
    "bodegas": Bodega,
    "plantas": Planta,
}


@app.get("/api/catalogos/{entidad}/uso")
def uso_catalogo(
    entidad: str, db: Session = Depends(obtener_sesion), _=Depends(usuario_actual)
):
    """Cuántos empleados y activos dependen de cada entrada del catálogo."""
    if entidad not in MODELOS_CATALOGO:
        raise HTTPException(404, "Catálogo desconocido")

    Modelo = MODELOS_CATALOGO[entidad]
    campo_emp = {
        "sucursales": Empleado.sucursal_id,
        "departamentos": Empleado.departamento_id,
        "bodegas": Empleado.bodega_id,
        "plantas": Empleado.planta_id,
    }[entidad]
    salida = []
    for registro in db.query(Modelo).order_by(Modelo.nombre).all():
        empleados = db.query(func.count(Empleado.id)).filter(campo_emp == registro.id).scalar()
        # Los activos no tienen ubicación propia: se cuentan por la persona que los tiene
        activos = (
            db.query(func.count(Asignacion.id))
            .join(Empleado, Empleado.id == Asignacion.empleado_id)
            .filter(campo_emp == registro.id, Asignacion.fecha_devolucion.is_(None))
            .scalar()
        )
        salida.append(
            {
                "id": registro.id,
                "nombre": registro.nombre,
                "clave": getattr(registro, "clave", None),
                "sucursal_id": getattr(registro, "sucursal_id", None),
                "sucursal": registro.sucursal.nombre
                if entidad == "bodegas" and registro.sucursal
                else None,
                "activo": registro.activo,
                "empleados": empleados,
                "activos": activos,
                "en_uso": empleados,
            }
        )
    return salida


@app.put("/api/catalogos/{entidad}/{registro_id}")
def editar_catalogo(
    entidad: str,
    registro_id: int,
    datos: schemas.CatalogoIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    if entidad not in MODELOS_CATALOGO:
        raise HTTPException(404, "Catálogo desconocido")
    registro = db.get(MODELOS_CATALOGO[entidad], registro_id)
    if registro is None:
        raise HTTPException(404, "Registro no encontrado")

    registro.nombre = datos.nombre.strip().upper()
    if entidad in ("sucursales", "departamentos"):
        if not datos.clave:
            raise HTTPException(400, "La clave es obligatoria")
        registro.clave = datos.clave.strip().upper()
    if entidad == "bodegas":
        registro.sucursal_id = datos.sucursal_id

    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Ya existe un registro con ese nombre o clave")
    registrar(db, usuario, "EDITAR", entidad, registro.id, registro.nombre)
    db.commit()
    return {"id": registro.id, "nombre": registro.nombre}


@app.delete("/api/catalogos/{entidad}/{registro_id}")
def eliminar_catalogo(
    entidad: str,
    registro_id: int,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    """Elimina una entrada del catálogo, sólo si nada depende de ella."""
    if entidad not in MODELOS_CATALOGO:
        raise HTTPException(404, "Catálogo desconocido")
    registro = db.get(MODELOS_CATALOGO[entidad], registro_id)
    if registro is None:
        raise HTTPException(404, "Registro no encontrado")

    campo_emp = {
        "sucursales": Empleado.sucursal_id,
        "departamentos": Empleado.departamento_id,
        "bodegas": Empleado.bodega_id,
        "plantas": Empleado.planta_id,
    }[entidad]
    empleados = db.query(func.count(Empleado.id)).filter(campo_emp == registro_id).scalar()
    if empleados:
        raise HTTPException(
            409,
            f"No se puede eliminar: {empleados} empleado(s) usan este registro. "
            "Reasígnalos primero o desactívalo.",
        )
    if entidad == "departamentos":
        impresoras = (
            db.query(func.count(Impresora.id))
            .filter(Impresora.departamentos.any(Departamento.id == registro_id))
            .scalar()
        )
        if impresoras:
            raise HTTPException(
                409, f"No se puede eliminar: {impresoras} impresora(s) usan este departamento"
            )
    campo_imp = {
        "sucursales": Impresora.sucursal_id,
        "bodegas": Impresora.bodega_id,
        "plantas": Impresora.planta_id,
    }.get(entidad)
    if campo_imp is not None:
        impresoras = db.query(func.count(Impresora.id)).filter(campo_imp == registro_id).scalar()
        if impresoras:
            raise HTTPException(
                409, f"No se puede eliminar: {impresoras} impresora(s) usan este registro"
            )
    if entidad == "sucursales":
        bodegas = db.query(func.count(Bodega.id)).filter(Bodega.sucursal_id == registro_id).scalar()
        if bodegas:
            raise HTTPException(409, f"No se puede eliminar: tiene {bodegas} bodega(s) asociada(s)")

    nombre = registro.nombre
    db.delete(registro)
    registrar(db, usuario, "ELIMINAR", entidad, registro_id, nombre)
    db.commit()
    return {"ok": True}


@app.post("/api/catalogos/{entidad}/{registro_id}/activo")
def alternar_catalogo(
    entidad: str,
    registro_id: int,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    """Activa o desactiva una entrada sin borrarla: deja de ofrecerse al capturar."""
    if entidad not in MODELOS_CATALOGO:
        raise HTTPException(404, "Catálogo desconocido")
    registro = db.get(MODELOS_CATALOGO[entidad], registro_id)
    if registro is None:
        raise HTTPException(404, "Registro no encontrado")
    registro.activo = not registro.activo
    registrar(
        db, usuario, "EDITAR", entidad, registro.id,
        f"{registro.nombre} {'activada' if registro.activo else 'desactivada'}",
    )
    db.commit()
    return {"id": registro.id, "activo": registro.activo}


# --------------------------------------------------------------------------
# Impresoras
# --------------------------------------------------------------------------


def dic_impresora(i: Impresora) -> dict:
    return {
        "id": i.id,
        "numero_serie": i.numero_serie,
        "propietario": i.propietario.value,
        "locacion": i.locacion,
        "departamentos": [{"id": d.id, "nombre": d.nombre} for d in i.departamentos],
        "departamento_ids": [d.id for d in i.departamentos],
        "sucursal_id": i.sucursal_id,
        "sucursal": i.sucursal.nombre if i.sucursal else None,
        "bodega_id": i.bodega_id,
        "bodega": i.bodega.nombre if i.bodega else None,
        "planta_id": i.planta_id,
        "planta": i.planta.nombre if i.planta else None,
        "numero_serie_proveedor": i.numero_serie_proveedor,
        "identificador_proveedor": i.identificador_proveedor,
    }


def _aplicar_impresora(db: Session, impresora: Impresora, datos: schemas.ImpresoraIn) -> None:
    ids = set(datos.departamento_ids)
    departamentos = db.query(Departamento).filter(Departamento.id.in_(ids)).all() if ids else []
    if len(departamentos) != len(ids):
        raise HTTPException(400, "Alguno de los departamentos no existe")
    impresora.numero_serie = datos.numero_serie.strip()
    impresora.propietario = datos.propietario
    impresora.locacion = (datos.locacion or "").strip() or None
    impresora.numero_serie_proveedor = (datos.numero_serie_proveedor or "").strip() or None
    impresora.identificador_proveedor = datos.identificador_proveedor.strip()
    impresora.departamentos = departamentos
    impresora.sucursal_id = datos.sucursal_id
    impresora.bodega_id = datos.bodega_id
    impresora.planta_id = datos.planta_id


@app.get("/api/impresoras")
def listar_impresoras(
    q: str | None = Query(None, max_length=100),
    propietario: PropietarioImpresora | None = None,
    departamento_id: int | None = None,
    sucursal_id: int | None = None,
    bodega_id: int | None = None,
    planta_id: int | None = None,
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    consulta = db.query(Impresora)
    if propietario:
        consulta = consulta.filter(Impresora.propietario == propietario)
    if departamento_id:
        consulta = consulta.filter(Impresora.departamentos.any(Departamento.id == departamento_id))
    if sucursal_id:
        consulta = consulta.filter(Impresora.sucursal_id == sucursal_id)
    if bodega_id:
        consulta = consulta.filter(Impresora.bodega_id == bodega_id)
    if planta_id:
        consulta = consulta.filter(Impresora.planta_id == planta_id)
    if q:
        patron = patron_busqueda(q)
        consulta = consulta.filter(
            or_(
                func.upper(Impresora.numero_serie).like(patron, escape="\\"),
                func.upper(Impresora.locacion).like(patron, escape="\\"),
                func.upper(Impresora.numero_serie_proveedor).like(patron, escape="\\"),
                func.upper(Impresora.identificador_proveedor).like(patron, escape="\\"),
            )
        )
    return [dic_impresora(i) for i in consulta.order_by(Impresora.numero_serie).all()]


@app.post("/api/impresoras")
def crear_impresora(
    datos: schemas.ImpresoraIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    impresora = Impresora()
    _aplicar_impresora(db, impresora, datos)
    db.add(impresora)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Ya existe una impresora con ese número de serie")
    registrar(db, usuario, "CREAR", "impresora", impresora.id, impresora.numero_serie)
    db.commit()
    return dic_impresora(impresora)


@app.put("/api/impresoras/{impresora_id}")
def editar_impresora(
    impresora_id: int,
    datos: schemas.ImpresoraIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    impresora = db.get(Impresora, impresora_id)
    if impresora is None:
        raise HTTPException(404, "Impresora no encontrada")
    _aplicar_impresora(db, impresora, datos)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Ya existe una impresora con ese número de serie")
    registrar(db, usuario, "EDITAR", "impresora", impresora.id, impresora.numero_serie)
    db.commit()
    return dic_impresora(impresora)


@app.delete("/api/impresoras/{impresora_id}")
def eliminar_impresora(
    impresora_id: int,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    impresora = db.get(Impresora, impresora_id)
    if impresora is None:
        raise HTTPException(404, "Impresora no encontrada")
    serie = impresora.numero_serie
    db.delete(impresora)
    registrar(db, usuario, "ELIMINAR", "impresora", impresora_id, serie)
    db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------
# Contraseñas de redes
# --------------------------------------------------------------------------


def dic_red(r: RedWifi) -> dict:
    return {
        "id": r.id,
        "alcance": r.alcance.value,
        "ubicacion": r.ubicacion,
        "ssid": r.ssid,
        "contrasena": descifrar(r.contrasena),
        "numero_cuenta": r.numero_cuenta,
        "proveedor": r.proveedor.value,
        "actualizado_en": r.actualizado_en,
    }


def _aplicar_red(red: RedWifi, datos: schemas.RedWifiIn) -> None:
    red.alcance = datos.alcance
    red.ubicacion = datos.ubicacion.strip()
    red.ssid = datos.ssid.strip()
    red.contrasena = cifrar(datos.contrasena)
    red.numero_cuenta = (datos.numero_cuenta or "").strip() or None
    red.proveedor = datos.proveedor


def _guardar_red(db: Session, usuario: Usuario, red: RedWifi, accion: str) -> dict:
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Esa red ya está registrada en esa ubicación")
    registrar(db, usuario, accion, "red", red.id, f"{red.ssid} · {red.ubicacion}")
    db.commit()
    return dic_red(red)


@app.get("/api/redes")
def listar_redes(
    alcance: AlcanceRed | None = None,
    db: Session = Depends(obtener_sesion),
    _=Depends(requiere_captura),
):
    consulta = db.query(RedWifi)
    if alcance:
        consulta = consulta.filter(RedWifi.alcance == alcance)
    return [dic_red(r) for r in consulta.order_by(RedWifi.ubicacion, RedWifi.ssid).all()]


@app.post("/api/redes")
def crear_red(
    datos: schemas.RedWifiIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_admin),
):
    red = RedWifi()
    _aplicar_red(red, datos)
    db.add(red)
    return _guardar_red(db, usuario, red, "CREAR")


@app.put("/api/redes/{red_id}")
def editar_red(
    red_id: int,
    datos: schemas.RedWifiIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_admin),
):
    red = db.get(RedWifi, red_id)
    if red is None:
        raise HTTPException(404, "Red no encontrada")
    _aplicar_red(red, datos)
    return _guardar_red(db, usuario, red, "EDITAR")


@app.delete("/api/redes/{red_id}")
def eliminar_red(
    red_id: int,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_admin),
):
    red = db.get(RedWifi, red_id)
    if red is None:
        raise HTTPException(404, "Red no encontrada")
    detalle = f"{red.ssid} · {red.ubicacion}"
    db.delete(red)
    registrar(db, usuario, "ELIMINAR", "red", red_id, detalle)
    db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------
# Notas y pendientes
# --------------------------------------------------------------------------


def dic_nota(n: Nota) -> dict:
    return {
        "id": n.id,
        "titulo": n.titulo,
        "contenido": n.contenido,
        "estado": n.estado.value,
        "prioridad": n.prioridad.value,
        "empleado_id": n.empleado_id,
        "empleado": n.empleado.nombre_completo if n.empleado else None,
        "activo_id": n.activo_id,
        "activo": n.activo.codigo if n.activo else None,
        "autor_nombre": n.autor_nombre,
        "resuelta_por": n.resuelta_por,
        "resuelta_en": n.resuelta_en,
        "creado_en": n.creado_en,
        "actualizado_en": n.actualizado_en,
    }


@app.get("/api/notas")
def listar_notas(
    estado: EstadoNota | None = None,
    empleado_id: int | None = None,
    activo_id: int | None = None,
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    consulta = db.query(Nota)
    if estado:
        consulta = consulta.filter(Nota.estado == estado)
    if empleado_id:
        consulta = consulta.filter(Nota.empleado_id == empleado_id)
    if activo_id:
        consulta = consulta.filter(Nota.activo_id == activo_id)

    orden_prioridad = {"ALTA": 0, "MEDIA": 1, "BAJA": 2}
    notas = consulta.order_by(Nota.creado_en.desc()).all()
    notas.sort(key=lambda n: (n.estado == EstadoNota.RESUELTA, orden_prioridad[n.prioridad.value]))
    return [dic_nota(n) for n in notas]


@app.post("/api/notas")
def crear_nota(
    datos: schemas.NotaIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    nota = Nota(
        **datos.model_dump(), autor_id=usuario.id, autor_nombre=usuario.nombre
    )
    db.add(nota)
    db.flush()
    registrar(db, usuario, "CREAR", "nota", nota.id, nota.titulo)
    db.commit()
    return dic_nota(nota)


@app.put("/api/notas/{nota_id}")
def editar_nota(
    nota_id: int,
    datos: schemas.NotaIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    nota = db.get(Nota, nota_id)
    if nota is None:
        raise HTTPException(404, "Nota no encontrada")

    era_abierta = nota.estado != EstadoNota.RESUELTA
    for campo, valor in datos.model_dump().items():
        setattr(nota, campo, valor)
    if era_abierta and nota.estado == EstadoNota.RESUELTA:
        nota.resuelta_por = usuario.nombre
        nota.resuelta_en = dt.datetime.now(dt.timezone.utc)
    elif nota.estado != EstadoNota.RESUELTA:
        nota.resuelta_por = None
        nota.resuelta_en = None

    registrar(db, usuario, "EDITAR", "nota", nota.id, nota.titulo)
    db.commit()
    return dic_nota(nota)


@app.delete("/api/notas/{nota_id}")
def eliminar_nota(
    nota_id: int,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    nota = db.get(Nota, nota_id)
    if nota is None:
        raise HTTPException(404, "Nota no encontrada")
    titulo = nota.titulo
    db.delete(nota)
    registrar(db, usuario, "ELIMINAR", "nota", nota_id, titulo)
    db.commit()
    return {"ok": True}



# --------------------------------------------------------------------------
# Verificación mensual de números
# --------------------------------------------------------------------------


def _solo_digitos(numero: str) -> str:
    digitos = re.sub(r"\D", "", numero or "")
    # Los números de México se guardan a 10 dígitos
    if len(digitos) == 12 and digitos.startswith("52"):
        digitos = digitos[2:]
    return digitos


def _sincronizar_lineas(db: Session) -> None:
    """Toda línea del inventario de telefonía entra a la verificación."""
    lineas = {l.numero: l for l in db.query(LineaVerificacion).all()}
    telefonos = (
        db.query(Activo)
        .join(DetalleMovil)
        .filter(Activo.tipo == TipoActivo.TELEFONO, DetalleMovil.numero_telefonico.isnot(None))
        .all()
    )
    for a in telefonos:
        numero = _solo_digitos(a.detalle_movil.numero_telefonico)
        if not numero:
            continue
        linea = lineas.get(numero)
        if linea is None:
            linea = LineaVerificacion(numero=numero, activo_id=a.id)
            db.add(linea)
            lineas[numero] = linea
        elif linea.activo_id is None:
            linea.activo_id = a.id
    if db.new or db.dirty:
        db.commit()


def dic_linea(l: LineaVerificacion, v: VerificacionNumero | None, db: Session) -> dict:
    titular, sucursal, departamento = l.titular, l.sucursal, l.departamento
    vig = asignacion_vigente(db, l.activo_id) if l.activo_id else None
    if vig:
        e = vig.empleado
        titular = e.nombre_completo
        sucursal = e.sucursal.nombre if e.sucursal else sucursal
        departamento = e.departamento.nombre if e.departamento else departamento
    return {
        "id": l.id,
        "numero": l.numero,
        "titular": titular,
        "titular_registrado": l.titular,
        "sucursal": sucursal,
        "sucursal_registrada": l.sucursal,
        "departamento": departamento,
        "departamento_registrado": l.departamento,
        "activo_id": l.activo_id,
        "activo_codigo": l.activo.codigo if l.activo else None,
        "empleado_id": vig.empleado_id if vig else None,
        "verificada": bool(v and v.verificada),
        "verificada_por": v.verificada_por if v else None,
        "verificada_en": v.verificada_en if v else None,
        "enviado_por": v.enviado_por if v else None,
        "enviado_en": v.enviado_en if v else None,
        "observacion": v.observacion if v else None,
    }


def _validar_periodo(periodo: str) -> str:
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", periodo):
        raise HTTPException(422, "Periodo inválido")
    return periodo


@app.get("/api/verificacion-numeros")
def listar_verificacion_numeros(
    periodo: str | None = None,
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    periodo = _validar_periodo(periodo or hoy_local().strftime("%Y-%m"))
    _sincronizar_lineas(db)
    verificaciones = {
        v.linea_id: v
        for v in db.query(VerificacionNumero).filter(VerificacionNumero.periodo == periodo)
    }
    lineas = (
        db.query(LineaVerificacion)
        .outerjoin(Activo, LineaVerificacion.activo_id == Activo.id)
        .filter(or_(Activo.id.is_(None), Activo.estatus != EstatusActivo.BAJA))
        .all()
    )
    filas = [dic_linea(l, verificaciones.get(l.id), db) for l in lineas]
    filas.sort(key=lambda f: ((f["titular"] or "~").strip().upper(), f["numero"]))
    return filas


def _guardar_linea(db: Session, linea: LineaVerificacion, datos: schemas.LineaVerificacionIn) -> None:
    numero = _solo_digitos(datos.numero)
    if len(numero) != 10:
        raise HTTPException(422, "El número debe tener 10 dígitos")
    otra = db.query(LineaVerificacion).filter(LineaVerificacion.numero == numero).first()
    if otra is not None and otra.id != linea.id:
        raise HTTPException(409, "Ese número ya está en la lista")
    linea.numero = numero
    linea.titular = (datos.titular or "").strip() or None
    linea.sucursal = (datos.sucursal or "").strip() or None
    linea.departamento = (datos.departamento or "").strip() or None


@app.post("/api/verificacion-numeros")
def crear_linea_verificacion(
    datos: schemas.LineaVerificacionIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    linea = LineaVerificacion()
    _guardar_linea(db, linea, datos)
    db.add(linea)
    db.flush()
    registrar(db, usuario, "CREAR", "verificacion", linea.id, linea.numero)
    db.commit()
    return {"id": linea.id}


@app.put("/api/verificacion-numeros/{linea_id}/{periodo}")
def marcar_verificacion(
    linea_id: int,
    periodo: str,
    datos: schemas.VerificacionIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    _validar_periodo(periodo)
    linea = db.get(LineaVerificacion, linea_id)
    if linea is None:
        raise HTTPException(404, "Número no encontrado")
    v = (
        db.query(VerificacionNumero)
        .filter(VerificacionNumero.linea_id == linea_id, VerificacionNumero.periodo == periodo)
        .first()
    )
    if v is None:
        v = VerificacionNumero(linea_id=linea_id, periodo=periodo, verificada=False)
        db.add(v)

    momento = dt.datetime.now(dt.timezone.utc)
    if datos.enviado:
        v.enviado_por = usuario.nombre
        v.enviado_en = momento
    if datos.verificada is not None and datos.verificada != v.verificada:
        v.verificada = datos.verificada
        v.verificada_por = usuario.nombre if datos.verificada else None
        v.verificada_en = momento if datos.verificada else None
        registrar(
            db, usuario, "EDITAR", "verificacion", linea.id,
            f"{linea.numero} {'verificado' if datos.verificada else 'sin verificar'} ({periodo})",
        )
    if "observacion" in datos.model_fields_set:
        v.observacion = (datos.observacion or "").strip() or None
    db.commit()
    return dic_linea(linea, v, db)


# --------------------------------------------------------------------------
# Accesos a sistemas internos
# --------------------------------------------------------------------------


@app.get("/api/empleados/{empleado_id}/accesos")
def ver_accesos(
    empleado_id: int, db: Session = Depends(obtener_sesion), _=Depends(usuario_actual)
):
    if db.get(Empleado, empleado_id) is None:
        raise HTTPException(404, "Empleado no encontrado")
    return accesos_de(db, empleado_id)


@app.put("/api/empleados/{empleado_id}/accesos")
def guardar_accesos(
    empleado_id: int,
    datos: schemas.AccesosEmpleadoIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    empleado = db.get(Empleado, empleado_id)
    if empleado is None:
        raise HTTPException(404, "Empleado no encontrado")

    existentes = {
        a.sistema_id: a
        for a in db.query(AccesoSistema).filter(AccesoSistema.empleado_id == empleado_id)
    }
    entregados = []
    for item in datos.accesos:
        if db.get(SistemaInterno, item.sistema_id) is None:
            raise HTTPException(404, f"Sistema {item.sistema_id} no encontrado")

        acceso = existentes.get(item.sistema_id)
        if acceso is None:
            acceso = AccesoSistema(empleado_id=empleado_id, sistema_id=item.sistema_id)
            db.add(acceso)

        cambio = acceso.entregado != item.entregado
        acceso.entregado = item.entregado
        acceso.usuario_sistema = item.usuario_sistema
        acceso.notas = item.notas
        # Si se marca entregado y no se puso fecha, queda la de hoy
        acceso.fecha_entrega = (
            item.fecha_entrega or (hoy_local() if item.entregado else None)
            if item.entregado
            else None
        )
        if cambio:
            acceso.actualizado_por = usuario.nombre
            acceso.actualizado_en = dt.datetime.now(dt.timezone.utc)
        if item.entregado:
            entregados.append(str(item.sistema_id))

    db.flush()
    registrar(
        db, usuario, "EDITAR", "accesos", empleado_id,
        f"{empleado.nombre_completo}: {len(entregados)} sistema(s) entregado(s)",
    )
    db.commit()
    return accesos_de(db, empleado_id)


@app.get("/api/sistemas")
def listar_sistemas(db: Session = Depends(obtener_sesion), _=Depends(usuario_actual)):
    salida = []
    for s in db.query(SistemaInterno).order_by(SistemaInterno.orden, SistemaInterno.nombre):
        entregados = (
            db.query(func.count(AccesoSistema.id))
            .filter(AccesoSistema.sistema_id == s.id, AccesoSistema.entregado.is_(True))
            .scalar()
        )
        salida.append(
            {
                "id": s.id,
                "nombre": s.nombre,
                "clave": s.clave,
                "descripcion": s.descripcion,
                "orden": s.orden,
                "activo": s.activo,
                "entregados": entregados,
            }
        )
    return salida


@app.post("/api/sistemas")
def crear_sistema(
    datos: schemas.SistemaIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    sistema = SistemaInterno(
        nombre=datos.nombre.strip().upper(),
        clave=datos.clave.strip().upper(),
        descripcion=datos.descripcion,
        orden=datos.orden,
    )
    db.add(sistema)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Ya existe un sistema con ese nombre o clave")
    registrar(db, usuario, "CREAR", "sistema", sistema.id, sistema.nombre)
    db.commit()
    return {"id": sistema.id, "nombre": sistema.nombre}


@app.put("/api/sistemas/{sistema_id}")
def editar_sistema(
    sistema_id: int,
    datos: schemas.SistemaIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    sistema = db.get(SistemaInterno, sistema_id)
    if sistema is None:
        raise HTTPException(404, "Sistema no encontrado")
    sistema.nombre = datos.nombre.strip().upper()
    sistema.clave = datos.clave.strip().upper()
    sistema.descripcion = datos.descripcion
    sistema.orden = datos.orden
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Ya existe un sistema con ese nombre o clave")
    registrar(db, usuario, "EDITAR", "sistema", sistema.id, sistema.nombre)
    db.commit()
    return {"id": sistema.id, "nombre": sistema.nombre}


@app.delete("/api/sistemas/{sistema_id}")
def eliminar_sistema(
    sistema_id: int,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    sistema = db.get(SistemaInterno, sistema_id)
    if sistema is None:
        raise HTTPException(404, "Sistema no encontrado")
    en_uso = (
        db.query(func.count(AccesoSistema.id))
        .filter(AccesoSistema.sistema_id == sistema_id, AccesoSistema.entregado.is_(True))
        .scalar()
    )
    if en_uso:
        raise HTTPException(
            409, f"No se puede eliminar: {en_uso} persona(s) tienen este acceso registrado."
        )
    nombre = sistema.nombre
    db.query(AccesoSistema).filter(AccesoSistema.sistema_id == sistema_id).delete()
    db.delete(sistema)
    registrar(db, usuario, "ELIMINAR", "sistema", sistema_id, nombre)
    db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------
# Auditoria semanal
# --------------------------------------------------------------------------


def dic_item(i: ItemAuditoria) -> dict:
    return {
        "id": i.id,
        "orden": i.orden,
        "categoria": i.categoria,
        "descripcion": i.descripcion,
        "automatico": i.automatico,
        "estado": i.estado.value,
        "observacion": i.observacion,
        "activo_id": i.activo_id,
        "activo": i.activo.codigo if i.activo else None,
        "actualizado_por": i.actualizado_por,
        "actualizado_en": i.actualizado_en,
        "evidencias": [dic_evidencia(e) for e in i.evidencias],
    }


def dic_evidencia(e: EvidenciaAuditoria) -> dict:
    return {
        "id": e.id,
        "item_id": e.item_id,
        "nombre": e.nombre,
        "url": e.url,
        "agregado_por": e.agregado_por,
        "agregado_en": e.agregado_en,
    }


def resumen_auditoria(a: Auditoria) -> dict:
    conteo = {e.value: 0 for e in EstadoItem}
    for i in a.items:
        conteo[i.estado.value] += 1
    total = len(a.items)
    revisados = total - conteo["PENDIENTE"]
    return {
        "id": a.id,
        "tipo": a.tipo.value,
        "fecha_inicio": a.fecha_inicio,
        "fecha_fin": a.fecha_fin,
        "estado": a.estado.value,
        "analista_id": a.analista_id,
        "analista": a.analista_nombre,
        "programada_por": a.programada_por,
        "observaciones": a.observaciones,
        "dias_sin_actividad": a.dias_sin_actividad,
        "fecha_cierre": a.fecha_cierre,
        "cerrada_por": a.cerrada_por,
        "total": total,
        "revisados": revisados,
        "avance": round(revisados / total * 100) if total else 0,
        "conteo": conteo,
        "incidencias": conteo["INCIDENCIA"],
    }


@app.get("/api/auditorias")
def listar_auditorias(
    tipo: TipoAuditoria | None = None,
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    consulta = db.query(Auditoria)
    if tipo:
        consulta = consulta.filter(Auditoria.tipo == tipo)
    return [
        resumen_auditoria(a)
        for a in consulta.order_by(Auditoria.fecha_inicio.desc(), Auditoria.tipo).all()
    ]


NOMBRE_AUDITORIA = {
    TipoAuditoria.CORREO: "Correo",
    TipoAuditoria.NORTON: "Norton",
}
ETIQUETA_ITEM = {
    EstadoItem.PENDIENTE: "Sin revisar",
    EstadoItem.CONFORME: "Conforme",
    EstadoItem.INCIDENCIA: "Incidencia",
    EstadoItem.NO_APLICA: "No aplica",
}
MESES = [
    "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
    "agosto", "septiembre", "octubre", "noviembre", "diciembre",
]


def _texto_fecha(valor) -> str:
    if isinstance(valor, dt.datetime):
        return f"{a_hora_local(valor):%d/%m/%Y %H:%M}"
    if isinstance(valor, dt.date):
        return f"{valor:%d/%m/%Y}"
    return ""


@app.get("/api/auditorias/reporte-mensual")
def reporte_mensual_auditoria(
    ids: list[int] = Query(..., min_length=1, max_length=6),
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    """Concentra en un Excel las auditorías semanales elegidas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    auditorias = (
        db.query(Auditoria)
        .filter(Auditoria.id.in_(ids))
        .order_by(Auditoria.fecha_inicio)
        .all()
    )
    if len(auditorias) != len(set(ids)):
        raise HTTPException(404, "Alguna de las auditorías elegidas no existe")
    tipos = {a.tipo for a in auditorias}
    if len(tipos) > 1:
        raise HTTPException(400, "Las auditorías elegidas deben ser del mismo tipo")
    tipo = tipos.pop()
    referencia = auditorias[len(auditorias) // 2].fecha_fin
    anio, numero = referencia.year, referencia.month
    mes = f"{anio}-{numero:02d}"

    titulo = f"Auditoría de {NOMBRE_AUDITORIA[tipo]} · {MESES[numero - 1]} {anio}"
    negrita = Font(bold=True)
    relleno = PatternFill("solid", fgColor="DDEBF7")
    colores = {
        EstadoItem.CONFORME: "C6EFCE",
        EstadoItem.INCIDENCIA: "FFC7CE",
        EstadoItem.PENDIENTE: "FFEB9C",
        EstadoItem.NO_APLICA: "EDEDED",
    }

    def encabezado(hoja, columnas: list[str], fila: int = 1) -> None:
        for i, texto in enumerate(columnas, start=1):
            celda = hoja.cell(row=fila, column=i, value=texto)
            celda.font = negrita
            celda.fill = relleno
            celda.alignment = Alignment(vertical="center", wrap_text=True)

    def anchos(hoja, medidas: list[int]) -> None:
        for i, ancho in enumerate(medidas, start=1):
            hoja.column_dimensions[get_column_letter(i)].width = ancho

    def ajustar(hoja, desde: int = 2) -> None:
        for fila in hoja.iter_rows(min_row=desde):
            for celda in fila:
                celda.alignment = Alignment(wrap_text=True, vertical="top")

    ordenados = {a.id: sorted(a.items, key=lambda x: (x.orden, x.id)) for a in auditorias}
    semana = {a.id: f"{a.fecha_inicio:%d/%m} — {a.fecha_fin:%d/%m/%Y}" for a in auditorias}
    wb = Workbook()

    hoja = wb.active
    hoja.title = "Resumen"
    hoja["A1"] = titulo
    hoja["A1"].font = Font(bold=True, size=14)
    encabezado(
        hoja,
        ["Semana", "Analista", "Estado", "Puntos", "Conformes", "Incidencias", "No aplica",
         "Sin revisar", "Avance", "Evidencias", "Cerrada por", "Fecha de cierre", "Observaciones"],
        3,
    )
    totales = [0] * 6
    for a in auditorias:
        r = resumen_auditoria(a)
        cifras = [
            r["total"], r["conteo"]["CONFORME"], r["conteo"]["INCIDENCIA"],
            r["conteo"]["NO_APLICA"], r["conteo"]["PENDIENTE"],
            sum(len(i.evidencias) for i in a.items),
        ]
        totales = [t + c for t, c in zip(totales, cifras)]
        hoja.append([
            semana[a.id], a.analista_nombre or "",
            "Cerrada" if a.estado == EstadoAuditoria.CERRADA else "En proceso",
            *cifras[:5], f"{r['avance']}%", cifras[5], a.cerrada_por or "",
            _texto_fecha(a.fecha_cierre), a.observaciones or "",
        ])
    cerradas = sum(a.estado == EstadoAuditoria.CERRADA for a in auditorias)
    revisados = totales[0] - totales[4]
    hoja.append([
        "Total del mes", "", f"{cerradas} de {len(auditorias)} cerradas", *totales[:5],
        f"{round(revisados / totales[0] * 100) if totales[0] else 0}%", totales[5],
    ])
    for celda in hoja[hoja.max_row]:
        celda.font = negrita
    ajustar(hoja, 4)
    anchos(hoja, [22, 24, 18, 9, 11, 12, 10, 12, 9, 11, 22, 18, 50])

    # Cada punto contra las semanas del mes; primero los de la plantilla
    hoja = wb.create_sheet("Concentrado")
    encabezado(hoja, ["Categoría", "Punto"] + [semana[a.id] for a in auditorias])
    filas: dict[tuple[str, str], list] = {}
    for indice, a in enumerate(auditorias):
        for i in ordenados[a.id]:
            filas.setdefault((i.categoria, i.descripcion), [None] * len(auditorias))[indice] = i.estado
    plantilla = {d for _, d in PLANTILLAS[tipo]}
    for (categoria, descripcion), estados in sorted(
        filas.items(), key=lambda f: f[0][1] not in plantilla and not f[0][1].startswith("Notificar")
    ):
        hoja.append([categoria, descripcion] + [ETIQUETA_ITEM[e] if e else "" for e in estados])
        for columna, e in enumerate(estados, start=3):
            if e:
                hoja.cell(row=hoja.max_row, column=columna).fill = PatternFill(
                    "solid", fgColor=colores[e]
                )
    ajustar(hoja)
    anchos(hoja, [24, 70] + [18] * len(auditorias))

    hoja = wb.create_sheet("Detalle")
    encabezado(hoja, ["Semana", "Categoría", "Punto", "Estado", "Observación",
                      "Revisó", "Fecha de revisión", "Evidencias"])
    for a in auditorias:
        for i in ordenados[a.id]:
            hoja.append([
                semana[a.id], i.categoria, i.descripcion, ETIQUETA_ITEM[i.estado],
                i.observacion or "", i.actualizado_por or "", _texto_fecha(i.actualizado_en),
                len(i.evidencias),
            ])
            hoja.cell(row=hoja.max_row, column=4).fill = PatternFill(
                "solid", fgColor=colores[i.estado]
            )
    ajustar(hoja)
    anchos(hoja, [22, 22, 60, 12, 45, 22, 18, 11])

    hoja = wb.create_sheet("Evidencias")
    encabezado(hoja, ["Semana", "Categoría", "Punto", "Archivo", "Enlace", "Agregó", "Fecha"])
    for a in auditorias:
        for i in ordenados[a.id]:
            for e in i.evidencias:
                hoja.append([
                    semana[a.id], i.categoria, i.descripcion, e.nombre, e.url,
                    e.agregado_por or "", _texto_fecha(e.agregado_en),
                ])
                enlace = hoja.cell(row=hoja.max_row, column=5)
                enlace.hyperlink = e.url
                enlace.font = Font(color="0563C1", underline="single")
    ajustar(hoja)
    anchos(hoja, [22, 22, 50, 32, 50, 22, 18])

    _neutralizar_formulas(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nombre = f"auditoria_{tipo.value.lower()}_{mes}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


@app.get("/api/auditorias/{auditoria_id}")
def ver_auditoria(
    auditoria_id: int, db: Session = Depends(obtener_sesion), _=Depends(usuario_actual)
):
    auditoria = db.get(Auditoria, auditoria_id)
    if auditoria is None:
        raise HTTPException(404, "Auditoría no encontrada")
    datos = resumen_auditoria(auditoria)
    datos["items"] = [
        dic_item(i) for i in sorted(auditoria.items, key=lambda x: (x.orden, x.id))
    ]
    return datos


@app.post("/api/auditorias")
def crear_auditoria(
    datos: schemas.AuditoriaIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_admin),
):
    # Primero los datos de entrada, después el conflicto de estado
    analista = db.get(Usuario, datos.analista_id)
    if analista is None or not analista.activo:
        raise HTTPException(404, "El analista indicado no existe o está desactivado")

    inicio, fin = semana_de(datos.fecha or hoy_local())
    existe = (
        db.query(Auditoria)
        .filter(Auditoria.tipo == datos.tipo, Auditoria.fecha_inicio == inicio)
        .first()
    )
    if existe:
        raise HTTPException(
            409,
            f"Ya existe la auditoría de {datos.tipo.value.lower()} para la semana del "
            f"{inicio:%d/%m/%Y}",
        )

    auditoria = Auditoria(
        tipo=datos.tipo,
        fecha_inicio=inicio,
        fecha_fin=fin,
        analista_id=analista.id,
        analista_nombre=analista.nombre,
        programada_por=usuario.nombre,
    )
    if datos.tipo == TipoAuditoria.NORTON:
        auditoria.dias_sin_actividad = datos.dias_sin_actividad or DIAS_SIN_ACTIVIDAD
    auditoria.items = construir_items(db, datos.tipo, inicio, fin, auditoria.dias_sin_actividad)
    db.add(auditoria)
    db.flush()
    registrar(
        db, usuario, "PROGRAMAR", "auditoria", auditoria.id,
        f"{datos.tipo.value} semana del {inicio:%d/%m/%Y} → {analista.nombre} "
        f"({len(auditoria.items)} puntos)",
    )
    db.commit()
    return ver_auditoria(auditoria.id, db)


@app.put("/api/auditorias/{auditoria_id}")
def editar_auditoria(
    auditoria_id: int,
    datos: schemas.AuditoriaEditarIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    auditoria = db.get(Auditoria, auditoria_id)
    if auditoria is None:
        raise HTTPException(404, "Auditoría no encontrada")

    if datos.estado == EstadoAuditoria.CERRADA and auditoria.estado != EstadoAuditoria.CERRADA:
        sin_revisar = [i for i in auditoria.items if i.estado == EstadoItem.PENDIENTE]
        if sin_revisar:
            raise HTTPException(
                409,
                f"Faltan {len(sin_revisar)} punto(s) por revisar. "
                "Márcalos como conforme, incidencia o no aplica antes de cerrar.",
            )
        auditoria.fecha_cierre = dt.datetime.now(dt.timezone.utc)
        auditoria.cerrada_por = usuario.nombre
    elif datos.estado == EstadoAuditoria.EN_PROCESO:
        auditoria.fecha_cierre = None
        auditoria.cerrada_por = None

    auditoria.estado = datos.estado
    auditoria.observaciones = datos.observaciones
    # Reasignar el analista es potestad del administrador
    if datos.analista_id and datos.analista_id != auditoria.analista_id:
        if usuario.rol != RolUsuario.ADMIN:
            raise HTTPException(403, "Sólo un administrador puede reasignar la auditoría")
        nuevo = db.get(Usuario, datos.analista_id)
        if nuevo is None or not nuevo.activo:
            raise HTTPException(404, "El analista indicado no existe o está desactivado")
        auditoria.analista_id = nuevo.id
        auditoria.analista_nombre = nuevo.nombre
    registrar(db, usuario, "EDITAR", "auditoria", auditoria.id, datos.estado.value)
    db.commit()
    return ver_auditoria(auditoria_id, db)


@app.put("/api/auditorias/items/{item_id}")
def actualizar_item(
    item_id: int,
    datos: schemas.ItemAuditoriaIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    item = db.get(ItemAuditoria, item_id)
    if item is None:
        raise HTTPException(404, "Punto no encontrado")
    if item.auditoria.estado == EstadoAuditoria.CERRADA:
        raise HTTPException(409, "La auditoría está cerrada. Reábrela para modificarla.")
    if datos.estado == EstadoItem.INCIDENCIA and not (datos.observacion or "").strip():
        raise HTTPException(400, "Una incidencia necesita que describas qué pasó")

    item.estado = datos.estado
    item.observacion = datos.observacion
    item.actualizado_por = usuario.nombre
    item.actualizado_en = dt.datetime.now(dt.timezone.utc)
    db.commit()
    return dic_item(item)


@app.post("/api/auditorias/{auditoria_id}/items")
def agregar_item(
    auditoria_id: int,
    datos: schemas.ItemExtraIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    auditoria = db.get(Auditoria, auditoria_id)
    if auditoria is None:
        raise HTTPException(404, "Auditoría no encontrada")
    if auditoria.estado == EstadoAuditoria.CERRADA:
        raise HTTPException(409, "La auditoría está cerrada")

    item = ItemAuditoria(
        auditoria_id=auditoria_id,
        orden=max((i.orden for i in auditoria.items), default=0) + 1,
        categoria=datos.categoria,
        descripcion=datos.descripcion,
    )
    db.add(item)
    db.flush()
    registrar(db, usuario, "CREAR", "item_auditoria", item.id, datos.descripcion[:80])
    db.commit()
    return dic_item(item)


DOMINIOS_DRIVE = ("drive.google.com", "docs.google.com")


@app.post("/api/auditorias/items/{item_id}/evidencias")
def agregar_evidencia(
    item_id: int,
    datos: schemas.EvidenciaIn,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    item = db.get(ItemAuditoria, item_id)
    if item is None:
        raise HTTPException(404, "Punto no encontrado")
    if item.auditoria.estado == EstadoAuditoria.CERRADA:
        raise HTTPException(409, "La auditoría está cerrada. Reábrela para modificarla.")
    if item.auditoria.tipo == TipoAuditoria.NORTON and item.estado != EstadoItem.INCIDENCIA:
        raise HTTPException(409, "Solo se agregan evidencias a los puntos con incidencia")
    url = datos.url.strip()
    partes = urlparse(url)
    if partes.scheme != "https" or partes.hostname not in DOMINIOS_DRIVE:
        raise HTTPException(400, "La evidencia debe ser un enlace de Google Drive")

    evidencia = EvidenciaAuditoria(
        item_id=item.id,
        nombre=(datos.nombre or "").strip() or "Archivo de Drive",
        url=url,
        agregado_por=usuario.nombre,
    )
    db.add(evidencia)
    db.flush()
    registrar(db, usuario, "CREAR", "evidencia_auditoria", evidencia.id, evidencia.nombre)
    db.commit()
    return dic_evidencia(evidencia)


@app.delete("/api/auditorias/evidencias/{evidencia_id}")
def eliminar_evidencia(
    evidencia_id: int,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_captura),
):
    evidencia = db.get(EvidenciaAuditoria, evidencia_id)
    if evidencia is None:
        raise HTTPException(404, "Evidencia no encontrada")
    if evidencia.item.auditoria.estado == EstadoAuditoria.CERRADA:
        raise HTTPException(409, "La auditoría está cerrada. Reábrela para modificarla.")
    nombre = evidencia.nombre
    db.delete(evidencia)
    registrar(db, usuario, "ELIMINAR", "evidencia_auditoria", evidencia_id, nombre)
    db.commit()
    return {"ok": True}


@app.delete("/api/auditorias/{auditoria_id}")
def eliminar_auditoria(
    auditoria_id: int,
    db: Session = Depends(obtener_sesion),
    usuario: Usuario = Depends(requiere_admin),
):
    auditoria = db.get(Auditoria, auditoria_id)
    if auditoria is None:
        raise HTTPException(404, "Auditoría no encontrada")
    etiqueta = f"{auditoria.tipo.value} semana del {auditoria.fecha_inicio:%d/%m/%Y}"
    db.delete(auditoria)
    registrar(db, usuario, "ELIMINAR", "auditoria", auditoria_id, etiqueta)
    db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------
# Datos para las graficas
# --------------------------------------------------------------------------


@app.get("/api/graficas/movimientos")
def grafica_movimientos(
    meses: int = Query(12, ge=3, le=36),
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    """Asignaciones y devoluciones por mes.

    Una reasignación no es un movimiento aparte: es la devolución de un
    activo más la asignación siguiente, y así aparece en la serie.
    """
    hoy = hoy_local()
    periodos: list[dict] = []
    anio, mes = hoy.year, hoy.month
    for _i in range(meses):
        periodos.append({"anio": anio, "mes": mes})
        mes -= 1
        if mes == 0:
            mes, anio = 12, anio - 1
    periodos.reverse()
    inicio = dt.date(periodos[0]["anio"], periodos[0]["mes"], 1)

    def por_mes(columna):
        filas = (
            db.query(
                func.strftime("%Y-%m", columna).label("periodo"),
                func.count(Asignacion.id),
            )
            .join(Activo, Activo.id == Asignacion.activo_id)
            .filter(columna.isnot(None), columna >= inicio, Activo.tipo.in_(TIPOS_FISICOS))
            .group_by("periodo")
            .all()
        )
        return {periodo: total for periodo, total in filas}

    asignaciones = por_mes(Asignacion.fecha_asignacion)
    devoluciones = por_mes(Asignacion.fecha_devolucion)

    NOMBRES = ["ene", "feb", "mar", "abr", "may", "jun",
               "jul", "ago", "sep", "oct", "nov", "dic"]
    salida = []
    for periodo in periodos:
        clave = f"{periodo['anio']:04d}-{periodo['mes']:02d}"
        salida.append(
            {
                "periodo": clave,
                "etiqueta": f"{NOMBRES[periodo['mes'] - 1]} {periodo['anio'] % 100:02d}",
                "asignaciones": asignaciones.get(clave, 0),
                "devoluciones": devoluciones.get(clave, 0),
            }
        )

    # Activos que han pasado por más de una persona
    reasignados = (
        db.query(func.count(Asignacion.activo_id))
        .group_by(Asignacion.activo_id)
        .having(func.count(Asignacion.id) > 1)
        .count()
    )
    return {
        "serie": salida,
        "total_asignaciones": sum(s["asignaciones"] for s in salida),
        "total_devoluciones": sum(s["devoluciones"] for s in salida),
        "activos_reasignados": reasignados,
    }


@app.get("/api/graficas/top-empleados")
def grafica_top_empleados(
    limite: int = Query(10, ge=3, le=30),
    db: Session = Depends(obtener_sesion),
    _=Depends(usuario_actual),
):
    """Quién concentra más activos vigentes, con el desglose por tipo."""
    filas = (
        db.query(Empleado, func.count(Asignacion.id).label("total"))
        .join(Asignacion, (Asignacion.empleado_id == Empleado.id))
        .join(Activo, Activo.id == Asignacion.activo_id)
        .filter(Asignacion.fecha_devolucion.is_(None), Activo.tipo.in_(TIPOS_FISICOS))
        .group_by(Empleado.id)
        .order_by(func.count(Asignacion.id).desc())
        .limit(limite)
        .all()
    )

    salida = []
    for empleado, total in filas:
        por_tipo = dict(
            db.query(Activo.tipo, func.count(Activo.id))
            .join(Asignacion, Asignacion.activo_id == Activo.id)
            .filter(
                Asignacion.empleado_id == empleado.id,
                Asignacion.fecha_devolucion.is_(None),
                Activo.tipo.in_(TIPOS_FISICOS),
            )
            .group_by(Activo.tipo)
            .all()
        )
        salida.append(
            {
                "empleado_id": empleado.id,
                "codigo": empleado.codigo,
                "nombre": empleado.nombre_completo,
                "departamento": empleado.departamento.nombre if empleado.departamento else None,
                "sucursal": empleado.sucursal.nombre if empleado.sucursal else None,
                "total": total,
                "por_tipo": {t.value: n for t, n in por_tipo.items()},
            }
        )
    return salida


# --------------------------------------------------------------------------
# Archivos estaticos
# --------------------------------------------------------------------------


@app.get("/")
def raiz():
    return FileResponse(ESTATICOS / "index.html")


app.mount("/", StaticFiles(directory=ESTATICOS), name="static")
